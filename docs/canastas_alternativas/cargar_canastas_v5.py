# ============================================================
# CARGAR LAS 6 CANASTAS EN "Productos unicos"  (script de Colab)  -- v5
# cantidad_01=Popular 02=Media 03=Ejecutiva 04=Tecnologica 05=Representativa 06=Femenina
# ------------------------------------------------------------
# GENERADO AUTOMATICAMENTE por docs/canastas_alternativas/construir_canastas_v5.py
#   periodo de referencia: 2026-08      generado: 2026-09-07
#   287 EANs   |   hogar de referencia: hogar tipo 2 (2 adultos + 2 ninos) = 3,09 adultos equivalentes
#
# Diferencia clave con v4: cada estrato usa SU PROPIA version de cada necesidad
# (Popular primer precio, Media marca lider, Ejecutiva premium), en vez de compartir
# los mismos EANs. Las cantidades salen de la cantidad FISICA mensual anclada a la
# CBA del INDEC dividida por el envase de cada producto, no de un conteo de unidades.
#
# Uso: pegar en Colab, ejecutar, subir el Excel canasta_representativa_*.xlsx
#      (cualquier version), descargar *_con_canastas.xlsx y subirlo a Drive en
#      carga/output_canasta/ como canasta_representativa_<periodo>.xlsx
#
# EDITABLE: se puede tocar cualquier cantidad a mano. Pero si se quiere recalibrar en
# serio (cambiar productos, cantidades fisicas o el hogar de referencia), conviene
# editar construir_canastas_v5.py y volver a generar este archivo, que deja registro
# de por que se eligio cada producto en canastas_v5_detalle.csv.
# ============================================================
import openpyxl, os
from google.colab import files

uploaded = files.upload()
archivo_excel = next(iter(uploaded.keys()))

SHEET = "Productos unicos"
COLS  = ["cantidad_01","cantidad_02","cantidad_03","cantidad_04","cantidad_05","cantidad_06"]

# ---- Cantidades por EAN (unidades de envase por mes) ----
CANTIDADES = {
    '198990785363': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Informática | Pc/Notebook Hp Note 15-Fc0235La R3-7320U Amd 15""""8/512Gb
    '34264476455': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Pequeños Electrodomésticos | Licuadora De Mano Oster 800W Fpsthb2802 354 Negro
    '4005808319695': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 1},  # Perfumería | Crema Corporal Milk Nutritiva Piel Extra Seca Nivea Body 4
    '4005808515592': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Crema Humectante con Aceite de Jojoba y Vitamina E Nivea S
    '4005900036759': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 10, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Desodorante Antitranspirante Hombre Roll On Nivea Power In
    '4005900985828': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1},  # Perfumería | Toallitas Desmaquillantes Micellair 3 en 1 Nivea 25 Un
    '4006000040363': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Crema Corporal Humectante con Aloe Vera 5 en 1 Nivea 400 M
    '53891142960': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Pequeños Electrodomésticos | Plancha A Vapor Oster Gcstbs5905
    '5900273001566': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Cepillo Dental Colgate Extra Clean 1 Un
    '650240035401': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 8, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Jabon Carbon Detox Asepxia 100 Gr
    '6932554416638': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Informática | Celular 14C 4Gb/128Gb Starry Blue Azul Redmi Xiaomi 1 Un
    '7500435178570': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Máquina de Afeitar Descartable Cuerpo Gillete 2 Un
    '7500435202671': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 7, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Shampoo Limpieza y Revitalización Head & Shoulders 180 Cc
    '7500435225366': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Repuesto Máquina Afeitar Carbono Gillette Mach 3 2 Un
    '7500435228763': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Bebés y mamás | Pañal Talle XG Baby-dry Hipoalergénico Pampers 36 Un
    '7500435228770': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebés y mamás | Pañal Talle XXG Baby-dry Hipoalergénico Pampers 34 Un
    '7500435229494': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Toallas Femeninas con Alas Noches Tranquilas Seca Always 1
    '7500435245852': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Máquina de Afeitar Desechable Carbón Prestobarba 3 Gillett
    '7500435247955': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Acondicionador Pro-V Miracles Biotinamina B3 Pantene 250 M
    '7502271278278': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Cocinas y Hornos | Cocina Multigas Cd5602Ab0 Drean 1 Un
    '7509546075242': {'cantidad_01': 0.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Cepillo Dental Super Flexi Suave Colgate 2 Un
    '7509546686523': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Perfumería | Crema Dental Triple Acción Colgate 90 Gr
    '7509552902815': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Acondicionador Keratin Rt5+ Elvive 400 Ml
    '7509552924121': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1},  # Perfumería | Tintura Coloración Permanente 50 Castaño Nutrisse 1 Un
    '7613287613356': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Mascotas | Alimento para Perros Adultos Bolsa Dogui 1.5 Kg
    '7622201702656': {'cantidad_01': 0, 'cantidad_02': 19, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Jugo en Polvo Sabor Limonada Clight 8 Gr
    '7622201703080': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 14, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Jugo en Polvo Sabor Manzana Deliciosa Clight 7 Gr
    '7622201705961': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 10, 'cantidad_06': 0},  # Bebidas | Jugo en Polvo Sabor Naranja Tang 15 Gr
    '7622201735685': {'cantidad_01': 13, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Jugo en Polvo Sabor Limón Dulce Tang 15 Gr
    '7622201735906': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 10, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Galletitas con Chips Chocolate Pepitos 119 Gr
    '7622201736033': {'cantidad_01': 0, 'cantidad_02': 6, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Galletitas Clásicas Cerealitas 212 Gr
    '7622300631574': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Chocolate con Relleno Oreo Milka 100 Gr
    '7622300829728': {'cantidad_01': 0, 'cantidad_02': 8, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Galletitas Chocolate con Relleno de Limón Melba 120 Gr
    '7702010631207': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Perfumería | Cepillo Dental Triple Accion Colgate 2 Un
    '7702018037865': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Perfumería | Repuesto Rasuradora Prestobarba Mach3 Sensitive 2 Un
    '7702018072392': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1},  # Perfumería | Rasuradora Desechable Simply Venus Gillette 2 Un
    '7702103751416': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Aritos de Cereal Frutados Froot Loops 195 Gr
    '7790010002653': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Perfumería | Toallas Femeninas con Alas Adapt Suave Siempre Libre 8 Un
    '7790010002769': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 4},  # Perfumería | Toallas Femeninas con Alas Ultrafina Suave Siempre Libre 8
    '7790010002806': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1},  # Perfumería | Protectores Diarios Protección Largos Carefree 60 Un
    '7790040133587': {'cantidad_01': 0, 'cantidad_02': 6, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Alfajor Minitorta Clásica Águila 69 Gr
    '7790040143234': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Almacén | Galletitas Chocolinas Paquete 262 Gr
    '7790040143364': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Galletitas Surtido Diversión 400 Gr
    '7790040144095': {'cantidad_01': 2.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Galletitas Traviata 540 Gr
    '7790040377806': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 13, 'cantidad_06': 0},  # Almacén | Galletitas Original Criollitas Pack 3 300 Gr
    '7790064000261': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1.5},  # Perfumería | Algodón Estrella Clásico 75 Gr
    '7790064001909': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Algodón Discos Redondos Estrella 80 Un
    '7790070012050': {'cantidad_01': 0, 'cantidad_02': 4, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Aceite de Girasol Cocinero 900 Ml
    '7790070034410': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Congelados | Nuggets de Pollo Bolsa Granja Del Sol 400 Gr
    '7790070035936': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Congelados | Bocaditos de Pollo Rebozados Patitas 400 Gr
    '7790070035950': {'cantidad_01': 0, 'cantidad_02': 0.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Congelados | Bocaditos de Pollo Rebozados +100 Gr Gratis Patitas 1.1 Kg
    '7790070036278': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Congelados | Espinaca Congelada Granja Del Sol 500 Gr
    '7790070036308': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Congelados | Brócoli Congelado Granja Del Sol 400 Gr
    '7790070231864': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Aceite de Oliva Cocinero 500 Cc
    '7790070320001': {'cantidad_01': 0, 'cantidad_02': 10, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Fideos Cabellos de Ángel Matarazzo 500 Gr
    '7790070335982': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 8, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Fideos Tirabuzón sin Gluten Matarazzo 500 Gr
    '7790070336118': {'cantidad_01': 12, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Fideos Tallarín N5 Lucchetti 500 Gr
    '7790070336385': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 11, 'cantidad_06': 0},  # Almacén | Fideos Spaghetti N7 Lucchetti 500 Gr
    '7790070418203': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 20, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Alfajor de Arroz Dulce de Leche Chocoarroz 22 Gr
    '7790070433169': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Pan Rallado Fortificado Preferido 500 Gr
    '7790070433176': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Pan Rallado Fortificado Preferido 1 Kg
    '7790070433251': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Rebozador Crunch Fortificado Preferido 450 Gr
    '7790070507280': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Yerba Mate Selección Nobleza Gaucha 1 Kg
    '7790070562258': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Almacén | Harina 000 con Vitamina Zinc Favorita 1 Kg
    '7790070562265': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Harina 0000 con Vitamina Zinc Favorita 1 Kg
    '7790070621863': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Ravioles Ricota Mozzarella Danbo Parmesano La Salteña 450 
    '7790070621870': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Frescos | Ravioles Jamón y Mozzarella La Salteña 450 Gr
    '7790070622037': {'cantidad_01': 1, 'cantidad_02': 1, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Tapa para Pascualina Criolla La Salteña 400 Gr
    '7790070622044': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Frescos | Tapa para Pascualina Hojaldradas La Salteña 400 Gr
    '7790072001014': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Sal Fina en Paquete Celusal 500 Gr
    '7790072001038': {'cantidad_01': 0, 'cantidad_02': 0.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Sal Fina en Salero Celusal 500 Gr
    '7790072002080': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Sal Fina en Estuche Celusal 500 Gr
    '7790088001497': {'cantidad_01': 5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Puré de Tomate en Tetrabrik Alco 520 Gr
    '7790117000200': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Limpieza | Bolsas de Residuos de 45 Cm 60 Cm en Rollo Asurin 30 Un
    '7790117061836': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Bolsas de Residuos Cierra Fácil Rollo 45x55 Cm Asurín 20 U
    '7790117061843': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Bolsas de Residuos Cierra Fácil Rollo 50x65 Cm Asurín 20 U
    '7790130000027': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Vinagre de Manzana Menoyo 500 Cc
    '7790130000034': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Vinagre de Alcohol Menoyo 500 Cc
    '7790130000058': {'cantidad_01': 0.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Vinagre de Alcohol Menoyo 1 Lt
    '7790132098459': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Limpieza | Lavandina Concentrada Ayudín 1 Lt
    '7790139003173': {'cantidad_01': 0, 'cantidad_02': 0.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Vinagre de Alcohol Clásico Casalta 1 Lt
    '7790150006320': {'cantidad_01': 0.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Café Molido Clásico Bolsa La Virginia 250 Gr
    '7790150465332': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Almacén | Comino Molido Sobre Alicante 25 Gr
    '7790150540183': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Orégano sin TACC Alicante 25 Gr
    '7790150564844': {'cantidad_01': 2.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Pimentón Dulce Sobre Alicante 25 Gr
    '7790150570180': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Provenzal Alicante 25 Gr
    '7790199000020': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Harina de Trigo 0000 Morixe 1 Kg
    '7790199604587': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Avena Tradicional Morixe 400 Gr
    '7790244001453': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Dulce de Leche Tradicional Artesanal San Ignacio 400 Gr
    '7790250015536': {'cantidad_01': 0, 'cantidad_02': 5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Papel Higiénico Doble Hoja Manzanilla Higienol 4 Un 30 Mt
    '7790250015840': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Limpieza | Papel Higiénico Hoja Simple Max Higienol 4 Un 100 Mt
    '7790250015857': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 6, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Papel Higiénico Hoja Simple Max Higienol 4 Un 80 Mt
    '7790250016182': {'cantidad_01': 4, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Papel Higiénico Fresh Hoja Simple con Aloe Vera Higienol 4
    '7790250022053': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Rollo de Cocina Floral Elite 60 Paños 3 Un
    '7790250056881': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Limpieza | Rollos de Cocina Clásico 50 Paños Sussex 3 Un
    '7790250057765': {'cantidad_01': 2, 'cantidad_02': 2.5, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Limpieza | Servilletas Papel Hoja Simple Clásica 30X30 Sussex 80 Un
    '7790272001005': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Aceite de Girasol Natura 900 Ml
    '7790272008394': {'cantidad_01': 0, 'cantidad_02': 0.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Aceite de Oliva Extra Vírgen Suave Natura 500 Ml
    '7790290101459': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Aperitivo Fernet Vittone 750 Cc
    '7790290101602': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Fernet Branca 750 Ml
    '7790310985366': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Papas Fritas Clásicas Pehuamar 135 Gr
    '7790310985458': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Almacén | Papas Fritas Clásicas Lays 85 Gr
    '7790310985465': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Papas Fritas Clásicas Lays 134 Gr
    '7790314080159': {'cantidad_01': 1.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Vino Malbec Estancia Mendoza 750 Cc
    '7790360967411': {'cantidad_01': 1.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Congelados | Hamburguesas Clásicas XL Swift 250 Gr
    '7790360967701': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Congelados | Hamburguesa Tipo Casera Swift 240 Gr
    '7790360970053': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Frescos | Salchichas Kids Swift Flowpack 190 Gr
    '7790380023401': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Polvo de Cacao Toddy 180 Gr
    '77903860': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 5, 'cantidad_06': 0},  # Almacén | Alfajor Triple Torta Terrabusi 70 Gr
    '7790387013627': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Yerba Mate con Palo 4Flex Taragui 500 Gr
    '7790398100088': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Manteca Calidad Extra La Paulina 200 Gr
    '7790520024954': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 6, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Limpiador Líquido Desinfectante Lavanda Doypack Lysoform 4
    '7790520025869': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Limpieza | Mata Mosca Mosquito Acuosa Aerosol Fuyi 360 Cc
    '7790520508607': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Insecticida Mata Moscas y Mosquitos sin Olor en Aerosol Ra
    '7790520997623': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Mata Moscas y Mosquitos Aerosol Raid 370 Ml
    '7790550000157': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Café Molido Torrado en Bolsa Cabrales 250 Gr
    '7790550022234': {'cantidad_01': 0, 'cantidad_02': 0.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Café Molido Tostado Súper Cabrales 250 Gr
    '7790580131357': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Atún en Aceite Libre de Gluten sin TACC La Campagnola 170 
    '7790580131364': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Almacén | Atún al Natural Libre de Gluten sin TACC La Campagnola 170
    '7790580132163': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Mermelada Durazno Arcor 454 Gr
    '7790580132392': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Almacén | Arvejas Secas Remojadas sin Conservantes Lata Arcor 300 Gr
    '7790580138721': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Polenta Prestopronta 730 Gr
    '7790580138738': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 1.5, 'cantidad_06': 0},  # Almacén | Harina de Maíz Paquete Prestopronta 500 Gr
    '7790580138868': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Almacén | Puré de Tomate La Campagnola 530 Gr
    '7790580567101': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Tomate Perita Cubeteado en Lata Arcor 400 Gr
    '7790580607210': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Almacén | Chocolate Leche Arcor 25 Gr
    '7790625010401': {'cantidad_01': 2.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Salchicha Viena 66 190 Gr
    '7790639003437': {'cantidad_01': 2.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Gaseosa Cola Classic Cunnington 2.25 Lt
    '7790639003468': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Gaseosa Lima Limón Classic Cunnington 2.25 Lt
    '7790670052586': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Congelados | Brocoli Congelado Green Life 450 Gr
    '7790742011701': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Crema de Leche Uat Tetra Top La Serenísima 200 Ml
    '7790742067005': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Dulce de Leche Repostero La Serenísima 400 Gr
    '7790742141101': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Frescos | Crema de Leche Uat Tetra Top La Serenísima 330 Ml
    '7790742192103': {'cantidad_01': 22, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Leche Descremada Sachet 1Casanto 1 Lt
    '7790742333605': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 29, 'cantidad_06': 0},  # Frescos | Leche UAT Parcialmente Descremada Zero Lactos Serenísima 1
    '7790742358608': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 31, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Leche Descremada Protein Larga Vida Uat La Serenisima 1 Lt
    '7790742363008': {'cantidad_01': 0, 'cantidad_02': 27, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Leche Larga Vida Entera Clásica 3% La Serenisima 1 Lt
    '7790742625304': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Frescos | Dulce de Leche Clásico La Serenísima 400 Gr
    '7790770601851': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Toallas Femeninas Nocturna Max Nosotras 8 Un
    '7790773007124': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Hisopos Family Q Soft 150 Un
    '7790895000447': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Gaseosa Sprite 1.5 Lt
    '7790895000829': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 7, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Gaseosa Sprite 500 Cc
    '7790895000997': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Coca Cola Sabor Original 2.25 Lt
    '7790895001000': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Bebidas | Gaseosa Sprite 2.25 Lt
    '7790895640483': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Bebidas | Agua Saborizada Manzana sin Gas Aquarius 1.5 Lt
    '7791120021558': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Arroz Doble Carolina Molinos Ala 1 Kg
    '7791120031557': {'cantidad_01': 4, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Arroz Largo Fino 00000 Molinos Ala 1 Kg
    '7791120031564': {'cantidad_01': 0, 'cantidad_02': 7, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Arroz Largo Fino 00000 Molinos Ala 500 Gr
    '7791120037559': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Almacén | Arroz Parboil Molinos Ala 1 Kg
    '7791120103858': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Polenta Instantánea Molinos Ala 500 Gr
    '7791130681384': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Quita Sarro Desinfectante en Gel Harpic 500 Ml
    '7791130683661': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Limpiador Líquido Marina Procenex 900 Cc
    '7791130963299': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Limpiador Líquido Baño Doypack Harpic 420 Ml
    '7791200000626': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Bebidas | Aperitivo Cynar 750 Ml
    '7791203001231': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Vino Tinto Malbec Luigi Bosca 750 Cc
    '7791290792050': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1.5, 'cantidad_06': 0},  # Limpieza | Jabón en Polvo Lavado a Mano Ala 800 Gr
    '7791290792142': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Jabón en Polvo Matic Mañana de Sol Ala 800 Gr
    '7791290792814': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Limpieza | Jabón Líquido para Ropa Bio Encimas Skip 800 Ml
    '7791290793583': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Limpieza | Suavizante para Ropa Concentrado Cuidado Esencial Doypack 
    '7791290793682': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Suavizante para Ropa Clásuci Doypack Comfort 900 Ml
    '7791290794054': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 7, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Detergente Líquido Bioactive Limón Cif 300 Ml
    '7791290794085': {'cantidad_01': 0, 'cantidad_02': 4, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Detergente Líquido Bioactive Limón Doypack Cif 450 Ml
    '7791290794115': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Limpieza | Detergente Líquido Bioactive Lima Cif 500 Ml
    '7791290794269': {'cantidad_01': 3, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Detergente Líquido Ultra Limón Doyoack Ala 450 Ml
    '7791290795112': {'cantidad_01': 2.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Jabón Líquido para Ropa Limón Doypack Granby 800 Ml
    '7791290795259': {'cantidad_01': 0, 'cantidad_02': 7, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Jabón Líquido para Ropa Fina Doypack Ala 450 Ml
    '7791290795600': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1.5, 'cantidad_06': 0},  # Limpieza | Limpiador Líquido Baño Expert Doypack Cif 450 Ml
    '7791290796409': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Jabón Liquido Limpieza Activo Skip 800 Ml
    '7791290796577': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza  | Jabón en Polvo mas Blancos Ala 600 Gr
    '7791293022581': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Desodorante Antitranspirante en Aerosol Rexona Extra Cool 
    '7791293043791': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Perfumería | Desodorante Aerosol Marine Axe 150 Cc
    '7791293045740': {'cantidad_01': 2.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Shampoo Crema Balance Sedal 340 Ml
    '7791293045948': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Acondicionador Crema Balance Sedal 340 Ml
    '7791293047102': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 1.5},  # Perfumería | Shampoo Reconstrucción Completa Dove 400 Ml
    '7791293047515': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1.5, 'cantidad_06': 1.5},  # Perfumería | Acondicionador Oleo Nutrición Dove 400 Cc
    '7791293049557': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 2.5},  # Perfumería | Desodorante Aerosol Women Nutrivive Rexona 150 Ml
    '7791293049595': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Desodorante Aerosol V8 Men Rexona 250 Ml
    '7791293050805': {'cantidad_01': 1.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Jabón Jazmín Cremoso Lux 360 Gr
    '7791293051017': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Shampoo Bond Intense Repair Dove 400 Ml
    '7791293051208': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 6, 'cantidad_06': 4},  # Perfumería | Jabón de Tocador Original Dove 90 Gr
    '7791337007246': {'cantidad_01': 0, 'cantidad_02': 13, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Yogur Firme Descremado con Colágeno Vainilla Pote Ser 190 
    '7791337008106': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 7, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Postre de Dulce de Leche con Crema Pote Danette 100 Gr
    '7791337008564': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 33, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Yogur Batido Parcialmente Descremado Vainilla con Calcio S
    '7791337008656': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Frescos | Postre Vainilla Pote Danonino 80.5 Gr
    '7791337008663': {'cantidad_01': 0, 'cantidad_02': 5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Postre Frutilla Pote Danonino 80.5 Gr
    '7791337008694': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 10, 'cantidad_06': 0},  # Frescos | Yogur Bebible Frutilla Danonino 185 Gr
    '7791337061422': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Frescos | Queso Untable Entero La Serenisima 290 Gr
    '7791337061439': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Queso Untable Light La Serenisima 290 Gr
    '7791564012488': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 6, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Queso Untable Jamón Adler 100 Gr
    '7791620187778': {'cantidad_01': 0.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Margarina Dorada Dánica 210 Gr
    '7791728248265': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Bebidas | Vino Tinto Dulce Santa Julia 750 Ml
    '7791813828464': {'cantidad_01': 0, 'cantidad_02': 4, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Gaseosa Cola Pepsi Black 2 Lt
    '7791813888468': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Bebidas | Gaseosa Cola Pepsi 2 Lt
    '7791828900377': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Rollo de Cocina Blanco 240 Paños Felpita 1 Un
    '7791866001364': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0.5, 'cantidad_06': 0},  # Almacén | Mayonesa Doypack Natura 500 Ml
    '7792180001641': {'cantidad_01': 4, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Aceite de Girasol Cañuelas 900 Ml
    '7792180004741': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Pan Rallado Fortificado Mamá Cocina 500 Gr
    '7792180139320': {'cantidad_01': 4, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Harina 000 Fortificada con Calcio Cañuelas 1 Kg
    '7792350067019': {'cantidad_01': 3, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Arvejas Secas Remojadas Inalpa 300 Gr
    '7792390620700': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Capelletini de Carne Giacomo 500 Gr
    '7792459163100': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Fibra Esponja Parrillera Go 1 Un
    '7792540250450': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2.5, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Almacén | Azúcar Molida Superior Ledesma 1 Kg
    '7792710000182': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Almacén | Yerba Mate sin TACC Amanda 500 Gr
    '7792798003716': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Bebidas | Cerveza en Botella No Retornable Corona 710 Cc
    '7792798007387': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Cerveza Rubia Cristal Quilmes 1 Lt
    '7792798010561': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Cerveza Blanca Botella Stella Artois 975 Cc
    '7792798014873': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 14, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Cerveza Corona 330 Ml
    '7793100111891': {'cantidad_01': 3, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Crema Dental Ultra Blanco Colgate 90 Gr
    '7793253003517': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Limpiador Desinfectante Líquido Lavanda Poett 1.8 Lt
    '7793253003548': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Limpieza | Limpiador Líquido Primavera Poett 900 Ml
    '7793253003722': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Lavandina Lavanda Triple Poder Ayudín 2 Lt
    '7793253003784': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Lavandina Anti-splash Tradicional Ayudín 2 Lt
    '7793253007102': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza  | Bolsas de Residuos 45X55 Cm de 15 Un Mortimer 1 Un
    '7793360000157': {'cantidad_01': 0, 'cantidad_02': 4, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Puré de Tomate Tetrabrik Salsati La Campagnola 520 Gr
    '7793360131516': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Mermelada de Durazno Bc La Campagnola 390 Gr
    '7793360131547': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Almacén | Mermelada de Ciruela Bc 390 Gr
    '7793360131561': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Mermelada Arándanos Bc La Campagnola 390 Gr
    '7793360132384': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Arvejas Secas Remojadas La Campagnola 300 Gr
    '7793433000497': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Sal Fina Modificada 66% Menos de Sodio Genser 300 Gr
    '7793704000928': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Yerba Suave Playadito 1 Kg
    '7793862008910': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Climatización | Convector Liliana Convectory Plus Mod Cfb17 1 Un
    '7793890258752': {'cantidad_01': 1.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Pan de Mesa Blanco en Rodajas Finas Lactal 460 Gr
    '7793890258769': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 10, 'cantidad_06': 0},  # Almacén | Pan de Mesa Blanco Chico Lactal 315 Gr
    '7793890258776': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Pan de Mesa Salvado Chico Lactal 330 Gr
    '7793890261486': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Pan Integral Bolsa Fargo 400 Gr
    '7793913013535': {'cantidad_01': 8, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Yogur Parcialmente Descremado Vainilla Tregar 125 Gr
    '7793940054006': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Frescos | Manteca La Serenísima 200 Gr
    '7794000006058': {'cantidad_01': 0, 'cantidad_02': 0.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Mayonesa Clásica Doypack Hellmanns 950 Gr
    '7794000007093': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Mayonesa Liviana sin TACC Doypack Hellmann's 237 Gr
    '7794000007109': {'cantidad_01': 0.5, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Mayonesa Liviana sin TACC Doypack Hellmann's 475 Gr
    '7794218106328': {'cantidad_01': 3, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Jabón en Pan con Glicerina Argentino 200 Gr
    '7794218106854': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 1.5, 'cantidad_04': 0, 'cantidad_05': 3, 'cantidad_06': 0},  # Limpieza | Jabón Blanco Argentino 150 Gr
    '7794440000098': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Esponja con Fibra Siempre Limpia Virulana 1 Un
    '7794440003303': {'cantidad_01': 2, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Esponja con Cuadraditos y Cuida Uñas Virulana 1 Un
    '7794440003341': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Limpieza | Esponja Multiuso Fibraesponjas Virulana Un
    '7794626011023': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Bebés y mamás | Toallitas Húmedas Clásico y Cotidiano Huggies 48 Un
    '7794626015298': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebés y mamás | Pañal XXXG Huggies Classic 28 Un
    '7794820902943': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 1.5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Dulce de Leche Repostero Milkaut 400 Gr
    '7794820903292': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Manteca Calidad Extra Milkaut 100 Gr
    '7794990878925': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Queso Untable Fundido con Jamón La Paulina 190 Gr
    '7795930000543': {'cantidad_01': 2.5, 'cantidad_02': 4, 'cantidad_03': 6, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Bebidas | Agua Mineral Neste Purez Vital 2.25 Lt
    '7796885495194': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Cocinas y Hornos | Hornos Microondas 20Lt Blanco 700 W Bgh 1 Un
    '7796962987321': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # CLIMATIZACIÓN | Aire Acondicionado Philco Split 3200W Frio Calor Cl. A Phs
    '7797026970037': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Informática | Parlante Bt Stromberg Boombox Force
    '7797102526875': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Heladeras y Freezers | Heladera con Freezer Cíclica 2 Fríos Modelo HDR280F50B Dre
    '7797453972291': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Mascotas | Alimento para Gatos Atún y Sardina Whiskas 1 Kg
    '7797470007426': {'cantidad_01': 0, 'cantidad_02': 4, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Arvejas Secas Remojadas Marolio 200 Gr
    '7798062548679': {'cantidad_01': 0, 'cantidad_02': 2.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Agua Saborizada Pomelo Sin Gas Levité 1.5 Lt
    '7798062548686': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 16, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Agua Saborizada Pomelo sin Gas Villa Del Sur Levité 500 Cc
    '7798074864675': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebidas | Vino Tinto Malbec Portillo 750 Ml
    '7798080520084': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Aceite de Oliva Extra Virgen Vidrio La Toscana 250 Ml
    '7798081287429': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Pequeños Electrodomésticos | Jarra Electrica Smart Life Mod Slek1714Wpn Blanca
    '7798131250380': {'cantidad_01': 3, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Lavandina Aditivada Naranja Héroe 1 Lt
    '7798137725820': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Informática | Auricular Ear Btwins 49 Noganet 1 Un
    '7798140257554': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1},  # Perfumería | Crema Facial Humectante Teatrical 100 Gr
    '7798174970016': {'cantidad_01': 0, 'cantidad_02': 4, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Azúcar Común Azucel 1 Kg
    '7798181510212': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 7, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Galletitas Crackers Smams 150 Gr
    '7798311610317': {'cantidad_01': 0, 'cantidad_02': 5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebés y mamás | Toallas Húmedas con Aloe Vera Toddler 48 Un
    '7798316700808': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 4, 'cantidad_06': 0},  # Almacén | Aceite de Girasol Legítimo 900 Ml
    '7798338291070': {'cantidad_01': 0, 'cantidad_02': 1, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Frescos | Crema de Leche Clásica para Cocinar Las Tres Niñas 200 Ml
    '7799086000389': {'cantidad_01': 4, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Azúcar Domino 1 Kg
    '77991584': {'cantidad_01': 4, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Alfajor Triple de Chocolate Fantoche 85 Gr
    '7891000345528': {'cantidad_01': 1, 'cantidad_02': 1, 'cantidad_03': 1, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Almacén | Caldo de Verdura Maggi 114 Gr
    '7891010251024': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 7, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Bebés y mamás | Toallitas Húmedas Johnsons 44 Un
    '7891010254773': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 2},  # Perfumería | Tampones Medio O.B. 8 Un
    '7891010560737': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 1, 'cantidad_06': 0},  # Perfumería | Hisopos Caja Johnson'S 75 Un
    '7891024116128': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Cepillo Dental Colgate 1 Un
    '7891024134429': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Crema Dental Sensitive Blanqueador Colgate 100 Gr
    '7891150000971': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 5, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Limpieza | Suavizante para Ropa Concentrado Cuidado Intesivo Comfor 5
    '7891150095618': {'cantidad_01': 0, 'cantidad_02': 7, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Jabón de Tocador Piel Sensible Dove 90 Gr
    '7891515629458': {'cantidad_01': 1, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Congelados | Formitas de Pollo Congelado Sadia 400 Gr
    '7896004004921': {'cantidad_01': 0, 'cantidad_02': 1.5, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Copos de Maíz Azucarados Zucaritas 240 Gr
    '7896004009155': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Papas Fritas Originales Pringles 104 Gr
    '7896015520045': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 4, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Perfumería | Crema Dental Blanqueadora Sensodyne Extra Fresh 90 Gr
    '8002990292139': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 1},  # Perfumería | Crema Depilatoria Piel Sensible Silck & Fresh Veet Caja 10
    '80052760': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Tableta de Chocolate Kinder Bueno 43 Gr
    '8445290057433': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2, 'cantidad_06': 0},  # Mascotas | Alimento para Gatos Adultos Pescado y Pollo Cat Chow 1 Kg
    '8445290068507': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 9, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Mascotas | Alimento para Gatos Carne Pollo Cat Chow 500 Gr
    '8445290274243': {'cantidad_01': 0, 'cantidad_02': 3, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Cacao en Polvo Nescao 150 Gr
    '8445290559975': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 2, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Cacao Menos Azúcar Nesquik 300 Gr
    '8445290944184': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 6, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Mascotas | Alimentos Animales Cachorro Mediano Grande Dog Chow 1.5 Kg
    '8445290977373': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Mascotas | Alimento para Perros Adultos Medianos a Grandes Bolsa Dog 
    '8445290987938': {'cantidad_01': 0, 'cantidad_02': 2, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Mascotas | Alimento para Perros Adultos Medianos y Pequeños Bolsa Dog
    '8445291082199': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 3, 'cantidad_04': 0, 'cantidad_05': 0, 'cantidad_06': 0},  # Almacén | Café Instantáneo Origen Dolca 100 Gr
    '8445291121843': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 0, 'cantidad_05': 2.5, 'cantidad_06': 0},  # Almacén | Polvo Chocolatado Nesquik 150 Gr
    '8718863045688': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Tv Audio y Video | Smart Tv 43" Full Hd 43Pfd6910/77 Philips 1 Un
    '8806094365573': {'cantidad_01': 0, 'cantidad_02': 0, 'cantidad_03': 0, 'cantidad_04': 1, 'cantidad_05': 0, 'cantidad_06': 0},  # Lavado | Lavarropas Frontal Inventer WW65A4000EEU Blanco 6.5 Kg Sam
}

wb = openpyxl.load_workbook(archivo_excel)
if SHEET not in wb.sheetnames:
    raise SystemExit(f"No encuentro la hoja '{SHEET}'. Hojas: {wb.sheetnames}")
ws = wb[SHEET]

hdr = {}
for j in range(1, ws.max_column + 1):
    v = ws.cell(row=1, column=j).value
    if v is not None:
        hdr[str(v).strip()] = j
for c in COLS:
    if c not in hdr:
        ws.cell(row=1, column=ws.max_column + 1, value=c)
        hdr[c] = ws.max_column

col_ean = None
for cand in ("id_producto", "ean", "ean_norm", "producto_sepa_id"):
    if cand in hdr:
        col_ean = hdr[cand]; break
if col_ean is None:
    raise SystemExit(f"No encuentro la columna de EAN. Columnas: {list(hdr)}")

# 1) limpiar: las 6 columnas se reescriben enteras
for i in range(2, ws.max_row + 1):
    for c in COLS:
        ws.cell(row=i, column=hdr[c], value=0)

# 2) cargar
puestos, encontrados = 0, set()
for i in range(2, ws.max_row + 1):
    raw = ws.cell(row=i, column=col_ean).value
    if raw is None:
        continue
    ean = str(raw).strip().lstrip("0")
    if ean in CANTIDADES:
        encontrados.add(ean)
        for c, q in CANTIDADES[ean].items():
            ws.cell(row=i, column=hdr[c], value=q)
        puestos += 1

faltan = sorted(set(CANTIDADES) - encontrados)
print(f"EANs de la canasta: {len(CANTIDADES)} | encontrados en la hoja: {len(encontrados)} | filas escritas: {puestos}")
if faltan:
    print(f"\nOJO - {len(faltan)} EANs NO estan en la hoja (no van a cotizar):")
    for e in faltan:
        print("   ", e, "|", " ".join(f"{k}={v}" for k, v in CANTIDADES[e].items() if v))
    print("\n   -> volve a correr construir_canastas_v5.py contra el Excel nuevo.")

for nombre, col in [("Popular","cantidad_01"),("Media","cantidad_02"),("Ejecutiva","cantidad_03"),
                    ("Tecnologica","cantidad_04"),("Representativa","cantidad_05"),("Femenina","cantidad_06")]:
    n = sum(1 for q in CANTIDADES.values() if q[col] > 0)
    u = sum(q[col] for q in CANTIDADES.values())
    print(f"  {nombre:15s} {n:3d} productos empaquetados | {u:8.1f} envases/mes")

salida = os.path.splitext(archivo_excel)[0] + "_con_canastas.xlsx"
wb.save(salida)
print(f"\nGuardado: {salida}")
files.download(salida)
