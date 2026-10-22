# ============================================================
# CANASTA DE PASS-THROUGH EN cantidad_01  (script de Colab)
# ------------------------------------------------------------
# Toma el Excel del mes canasta_representativa_AAAA-MM.xlsx y devuelve una COPIA con la canasta fija de pass-through
# (775 productos con precio todos los días del 01/01/2022 al 20/10/2026; ver docs/canastas_alternativas/README.md) en la
# columna cantidad_01 de la hoja Selección, que es la que lee el nb08 por defecto.
#
# Uso: pegar en una celda de Colab y ejecutar. Pide subir archivos:
#   - el Excel del mes (canasta_representativa_AAAA-MM.xlsx), y
#   - la canasta (canasta_pass_through_20220101_20261020.xlsx o .csv), salvo que esté en GitHub (URL_CANASTA).
# Devuelve canasta_pass_through_AAAA-MM.xlsx (se descarga y, si GUARDAR_EN_DRIVE, se copia al Drive).
#
# - Los productos de la canasta que no están en Selección se agregan copiando su fila de "Productos unicos" (sin sus
#   cantidades, que son del nb07) o, si no está, con los datos de la canasta.
# - cantidad_01 queda solo con la canasta; si tenía cantidades, se guardan en cantidad_01_original. Las demás
#   cantidad_0X y las otras hojas no se tocan.
# - La salida NO se llama canasta_representativa_*: el nb02 y el nb07 toman el canasta_representativa_*.xlsx de nombre
#   más alto de output_canasta y la leerían como si fuera la del mes.
# ============================================================

# ---------------- CONFIGURACIÓN ----------------
CANTIDAD = 'ipc'      # 'ipc'   = pesos del IPC por grupo (la canasta vale $1.000.000 a precios medianos de oct-2026)
                      # 'igual' = cada grupo del IPC pesa lo mismo
                      # 'uno'   = 1 unidad de cada producto (solo marca qué productos entran)
TIPOS = None          # None = los 775 productos. Para marcar una parte, una lista con alguno de: 'Alimentos',
                      # 'Bebidas sin alcohol', 'Bebidas alcohólicas', 'Limpieza del hogar',
                      # 'Cuidado personal y botiquín', 'Mascotas'
GUARDAR_EN_DRIVE = True
CARPETA_DRIVE = '/content/drive/MyDrive/carga/output_canasta'
URL_CANASTA = 'https://raw.githubusercontent.com/santiagoriverti/precios_minoristas_supermercados/main/docs/canastas_alternativas/canasta_pass_through_20220101_20261020.csv'
# -----------------------------------------------

import os, re, shutil, unicodedata
import pandas as pd
import openpyxl
from google.colab import files

TIPOS_VALIDOS = ['Alimentos', 'Bebidas sin alcohol', 'Bebidas alcohólicas', 'Limpieza del hogar',
                 'Cuidado personal y botiquín', 'Mascotas']

def _sin_tildes(s):
    return unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower().strip()

def _ean(v):
    return '' if v is None else str(v).strip().split('.')[0].lstrip('0')

def _num(v):
    try:
        return float(str(v).replace(',', '.')) if v not in (None, '') else None
    except ValueError:
        return None

print('Subí el Excel del mes (canasta_representativa_AAAA-MM.xlsx) y, si la canasta no está en GitHub, también '
      'canasta_pass_through_20220101_20261020.xlsx (se pueden elegir los dos juntos).')
subidos = list(files.upload().keys())

def es_canasta(n):
    # la canasta es el CSV o el Excel que entregó la selección (tiene la hoja Grupos_IPC), no por el nombre: la salida
    # de este script también se llama canasta_pass_through_*
    if n.lower().endswith('.csv'):
        return True
    try:
        return 'Grupos_IPC' in openpyxl.load_workbook(n, read_only=True).sheetnames
    except Exception:
        return False

archivos_pt = [n for n in subidos if es_canasta(n)]
archivos_mes = [n for n in subidos if n not in archivos_pt and n.lower().endswith(('.xlsx', '.xlsm'))]
if not archivos_mes:
    raise SystemExit(f'No subiste el Excel del mes (subiste: {subidos})')
archivo_mes = archivos_mes[0]
if len(archivos_mes) > 1:
    print(f'OJO: subiste {len(archivos_mes)} Excel; uso {archivo_mes}')

# ---- La canasta ----
if archivos_pt:
    n = archivos_pt[0]
    canasta = (pd.read_csv(n, dtype={'id_producto': str}) if n.lower().endswith('.csv')
               else pd.read_excel(n, sheet_name='Selección', dtype={'id_producto': str}))
    fuente = n
else:
    try:
        canasta = pd.read_csv(URL_CANASTA, dtype={'id_producto': str})
        fuente = 'GitHub'
    except Exception as e:
        raise SystemExit(f'No pude bajar la canasta de GitHub ({e}).\nVolvé a ejecutar y subí también '
                         'canasta_pass_through_20220101_20261020.xlsx junto con el Excel del mes.')
canasta['ean'] = canasta['id_producto'].map(_ean)
if CANTIDAD == 'ipc':
    canasta['q'] = canasta['cantidad_08']
elif CANTIDAD == 'igual':
    canasta['q'] = canasta['peso_igual_grupo_%'] / 100 * 1_000_000 / canasta['precio_mediano_oct26']
elif CANTIDAD == 'uno':
    canasta['q'] = 1.0
else:
    raise SystemExit(f"CANTIDAD tiene que ser 'ipc', 'igual' o 'uno' (puse {CANTIDAD!r})")
canasta['q'] = pd.to_numeric(canasta['q'], errors='coerce').round(6)
if TIPOS:
    _mal = [t for t in TIPOS if t not in TIPOS_VALIDOS]
    if _mal:
        raise SystemExit(f'Tipos que no existen: {_mal}. Los tipos son: {TIPOS_VALIDOS}')
    canasta = canasta[canasta['tipo'].isin(TIPOS)]
canasta = canasta[canasta['q'] > 0].drop_duplicates('ean')
Q = dict(zip(canasta['ean'], canasta['q']))
INFO = canasta.set_index('ean')
print(f'\nCanasta ({fuente}): {len(Q)} productos | cantidades: {CANTIDAD}' + (f' | tipos: {TIPOS}' if TIPOS else ''))

# ---- El Excel del mes ----
print(f'Abriendo {archivo_mes} (un Excel de ~10 MB tarda un minuto)...')
wb = openpyxl.load_workbook(archivo_mes)
hoja = next((s for s in wb.sheetnames if _sin_tildes(s) == 'seleccion'), None)
creada = hoja is None
if creada:
    print(f'OJO: {archivo_mes} no tiene la hoja Selección (¿es el canasta_analisis del nb02?). Se crea una con la canasta.')
    hoja = 'Selección'
    ws = wb.create_sheet(hoja)
    for j, c in enumerate(['id_producto', 'descripcion', 'marca', 'presentacion', 'rubro', 'categoria', 'subcategoria',
                           'cantidad_01'], start=1):
        ws.cell(row=1, column=j, value=c)
ws = wb[hoja]

hdr = {}
for j in range(1, ws.max_column + 1):
    v = ws.cell(row=1, column=j).value
    if v is not None:
        hdr[str(v).strip()] = j

def columna(nombre):
    if nombre not in hdr:
        hdr[nombre] = ws.max_column + 1
        ws.cell(row=1, column=hdr[nombre], value=nombre)
    return hdr[nombre]

col_id = next((hdr[c] for c in ('id_producto', 'ean', 'ean_norm', 'producto_sepa_id') if c in hdr), None)
if col_id is None:
    raise SystemExit(f'La hoja {hoja} no tiene columna de EAN (id_producto). Columnas: {list(hdr)}')
n_filas = ws.max_row
c01 = columna('cantidad_01')
originales = [ws.cell(row=i, column=c01).value for i in range(2, n_filas + 1)]
if any(_num(v) for v in originales) and 'cantidad_01_original' not in hdr:
    cbak = columna('cantidad_01_original')
    for i, v in enumerate(originales, start=2):
        ws.cell(row=i, column=cbak, value=v)
    print('cantidad_01 tenía cantidades: quedan guardadas en cantidad_01_original')
cflag, cgrupo, ctipo = columna('pass_through'), columna('pt_grupo_ipc'), columna('pt_tipo')

def marcar(i, e):
    ws.cell(row=i, column=c01, value=float(Q[e]))
    ws.cell(row=i, column=cflag, value=1)
    ws.cell(row=i, column=cgrupo, value=f"{INFO.at[e, 'grupo_ipc_cod']} {INFO.at[e, 'grupo_ipc']}")
    ws.cell(row=i, column=ctipo, value=INFO.at[e, 'tipo'])

vistos = set()
for i in range(2, n_filas + 1):
    e = _ean(ws.cell(row=i, column=col_id).value)
    if e in Q and e not in vistos:
        marcar(i, e)
        vistos.add(e)
    else:
        for c in (c01, cflag, cgrupo, ctipo):
            ws.cell(row=i, column=c).value = None      # ws.cell(..., value=None) no borra la celda

# ---- Productos de la canasta que no están en Selección: se agregan al final ----
faltan = [e for e in Q if e not in vistos]
de_unicos = {}
hu = next((s for s in wb.sheetnames if _sin_tildes(s) == 'productos unicos'), None)
if faltan and hu:
    wu = wb[hu]
    hdr_u = {str(wu.cell(row=1, column=j).value).strip(): j for j in range(1, wu.max_column + 1)
             if wu.cell(row=1, column=j).value is not None}
    cid_u = next((hdr_u[c] for c in ('id_producto', 'ean', 'ean_norm') if c in hdr_u), None)
    if cid_u:
        buscar = set(faltan)
        for fila in wu.iter_rows(min_row=2, values_only=True):
            e = _ean(fila[cid_u - 1])
            if e in buscar and e not in de_unicos:
                de_unicos[e] = {h: fila[j - 1] for h, j in hdr_u.items()}
periodo_hoja = next((ws.cell(row=i, column=hdr['periodo']).value for i in range(2, min(n_filas, 50) + 1)
                     if 'periodo' in hdr and ws.cell(row=i, column=hdr['periodo']).value), None)
DE_CANASTA = {'id_producto': 'id_producto', 'descripcion': 'descripcion', 'marca': 'marca', 'presentacion': 'presentacion',
              'rubro': 'rubro', 'categoria': 'categoria', 'subcategoria': 'subcategoria'}
i = n_filas
for e in faltan:
    i += 1
    datos = de_unicos.get(e)
    for h, j in hdr.items():
        if h.startswith('cantidad') or h in ('pass_through', 'pt_grupo_ipc', 'pt_tipo'):
            continue                  # las cantidades de "Productos unicos" son las del nb07: no se copian
        if datos is not None:
            v = datos.get(h)
        elif h in DE_CANASTA:
            v = INFO.at[e, DE_CANASTA[h]] if DE_CANASTA[h] in INFO.columns else None
        elif h == 'periodo':
            v = periodo_hoja
        else:
            v = None
        if h == 'id_producto' or j == col_id:
            v = str(e)
        if v is not None and not (isinstance(v, float) and pd.isna(v)):
            ws.cell(row=i, column=j, value=v)
    marcar(i, e)

# ---- Guardar ----
_m = re.search(r'(\d{4}-\d{2})', str(periodo_hoja or '')) or re.search(r'(\d{4}-\d{2})', archivo_mes)
salida = f"canasta_pass_through_{_m.group(1) if _m else 'sin_periodo'}.xlsx"
print('Guardando...')
wb.save(salida)

valor = (canasta['q'] * canasta['precio_mediano_oct26']).sum() if 'precio_mediano_oct26' in canasta else float('nan')
print(f'\nHoja {hoja}{" (creada)" if creada else ""}: {len(vistos)} productos de la canasta ya estaban | '
      f'{len(faltan)} agregados al final ({len(de_unicos)} copiados de "Productos unicos")')
print(f'cantidad_01 = canasta de pass-through: {len(Q)} productos | valor a precios medianos de oct-2026: ${valor:,.0f}')
print('Por tipo:', canasta['tipo'].value_counts().to_dict())
print(f'\nGuardado: {salida}')
if GUARDAR_EN_DRIVE:
    from google.colab import drive
    drive.mount('/content/drive')
    os.makedirs(CARPETA_DRIVE, exist_ok=True)
    shutil.copy(salida, CARPETA_DRIVE)
    print(f'Copiado a {CARPETA_DRIVE}/{salida}')
print(f"""
Para el nb08 (valor diario de la canasta), en la CELDA 1:
    CANASTA_EXCEL = '{CARPETA_DRIVE}/{salida}'
    CANASTA_COL   = 'cantidad_01'
El nb08 llama "Vulnerable" a cantidad_01 (en títulos y en el nombre del Excel de salida). Con los 775 productos la
primera corrida lleva unas 3 horas en el Colab gratuito (~5 GB de RAM por mes; guarda cada mes en caché y, si se
corta, sigue donde quedó). Para probar antes, TIPOS = ['Alimentos'].""")
files.download(salida)
