# Excel de decisión para elegir los productos de la brecha celíaca: pares de la misma marca, sin TACC multicadena,
# historia 2022-2026 (nb09) y ponderaciones de la ENGHo 2017/18.
# Uso: python -I -X utf8 s6_consolidar_excel.py <salida.xlsx de s4> <salida .parquet de s5>
#                                     <historia_serie.parquet> <engho2018_gastos.zip> <engho2018_articulos.zip> <salida.xlsx>
import sys, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

F_CAND, F_HIST, F_SER, F_GAS, F_ART, F_OUT = (Path(a) for a in sys.argv[1:7])
norm = lambda s: s.astype(str).str.strip().str.lstrip('0')

H = pd.read_parquet(F_HIST); H['id_producto'] = norm(H['id_producto']); Hi = H.set_index('id_producto')
SER = pd.read_parquet(F_SER)
PR = pd.read_parquet(F_CAND.with_suffix('.pares.parquet'))
for c in ['sin_id', 'con_id']:
    PR[c] = norm(PR[c])
REF = pd.read_excel(F_CAND, sheet_name='Referencias_con_TACC', dtype={'id_producto': str})
TODOS = pd.read_excel(F_CAND, sheet_name='Todos_por_tipo', dtype={'id_producto': str})

# ── Pares recomendados (revisados a mano: misma marca, mismo producto, presentación parecida) ───────────────────────
PARES = [
    ('Fideos secos', 'Spaghetti 500 g', '7790070336002', '7790070336309', '$/100g', ''),
    ('Fideos secos', 'Penne / mostachol 500 g', '7790070335999', '7790070336217', '$/100g', ''),
    ('Fideos secos', 'Tirabuzón 500 g', '7790070335982', '7790070336293', '$/100g',
     'sin TACC con cobertura irregular desde 2024-05 (mínimo 26% de su mediana)'),
    ('Premezclas (bizcochuelo)', 'Bizcochuelo vainilla Exquisita', '7790070432933', '7790070432964', '$/100g',
     'sin TACC 450 g vs 540 g'),
    ('Premezclas (bizcochuelo)', 'Bizcochuelo chocolate Exquisita', '7790070432940', '7790070432971', '$/100g',
     'sin TACC 450 g vs 540 g; mínimo 31% de su mediana desde 2024-05'),
    ('Premezclas (brownie)', 'Brownie Exquisita', '7790070432957', '7790070432858', '$/100g', '400 g vs 425 g'),
    ('Tapas', 'Tapas empanadas La Salteña x12', '7790070622075', '7790070622020', '$/paquete',
     'comparar por paquete de 12 (los gramos no vienen en todos los comercios)'),
    ('Tapas', 'Tapas pascualina / tarta La Salteña x2', '7790070622082', '7790070622037', '$/paquete',
     'pascualina sin TACC vs tarta criolla'),
    ('Galletitas dulces', 'Oreo', '7622202038099', '7622201735296', '$/100g',
     'sin TACC 95 g vs 118 g; sin TACC estable recién desde 2024-07 (en 2024-05 tenía 379 sucursales/día)'),
    ('Pan rallado', 'Pan rallado Preferido', '7790070100283', '7790070433169', '$/100g',
     'sin TACC 350 g vs 500 g; sin TACC desde 2025-04'),
    ('Alfajores', 'Alfajor Terrabusi', '7622202296291', '77903792', '$/100g',
     'sin gluten 55 g vs triple 70 g; sin gluten desde 2025-07'),
    ('Premezclas (sustitución)', 'Premezcla pizza Blancaflor vs harina leudante', '7790070508010', '7790070562180',
     '$/100g', 'SUSTITUCIÓN, no el mismo producto; sin TACC sin dato en dic-23/abr-24'),
]
TD = TODOS.copy(); TD['id_producto'] = norm(TD['id_producto']); TD = TD.drop_duplicates('id_producto').set_index('id_producto')
rows = []
for tipo, nombre, s, c, unidad, nota in PARES:
    a, b = TD.loc[s], TD.loc[c]
    ps, pc = (a['precio_100g'], b['precio_100g']) if unidad == '$/100g' else (a['precio_mediano'], b['precio_mediano'])
    hs, hc = Hi.loc[s], Hi.loc[c]
    rows.append({
        'tipo': tipo, 'par': nombre,
        'sin_TACC_ean': s, 'sin_TACC_descripcion': a['descripcion'], 'sin_TACC_gramos': a['gramos'],
        'sin_TACC_sucursales_22oct': a['sucursales'], 'sin_TACC_cadenas_22oct': a['cadenas'],
        'con_TACC_ean': c, 'con_TACC_descripcion': b['descripcion'], 'con_TACC_gramos': b['gramos'],
        'con_TACC_sucursales_22oct': b['sucursales'], 'con_TACC_cadenas_22oct': b['cadenas'],
        'comparacion': unidad, 'precio_sin_TACC': round(ps, 1), 'precio_con_TACC': round(pc, 1),
        'brecha_22oct_%': round((ps / pc - 1) * 100, 1),
        'sin_TACC_desde': hs['primer_mes'], 'con_TACC_desde': hc['primer_mes'],
        'sin_TACC_estable_desde_2024_05': bool(hs['estable_desde_2024_05']),
        'con_TACC_estable_desde_2024_05': bool(hc['estable_desde_2024_05']),
        'sin_TACC_min_sobre_mediana_2024_05': hs['min_sobre_mediana_2024_05'], 'nota': nota})
REC = pd.DataFrame(rows)

# ── Ponderaciones ENGHo 2017/18 (gasto de los hogares por artículo; hogar promedio, no celíaco) ─────────────────────
TIPO_ENGHO = {
    'A0111304': 'Fideos secos', 'A0111105': 'Galletitas dulces', 'A0111108': 'Galletitas dulces',
    'A0111103': 'Galletitas saladas, crackers y tostadas', 'A0111102': 'Galletitas saladas, crackers y tostadas',
    'A0111107': 'Galletitas saladas, crackers y tostadas', 'A0111114': 'Pan rallado y rebozador',
    'A0111109': 'Pan envasado', 'A0111110': 'Pan envasado', 'A0111113': 'Pan envasado',
    'A0111119': 'Budines, bizcochuelos y tortas', 'A0119205': 'Premezclas y harinas', 'A0111210': 'Premezclas y harinas',
    'A0111211': 'Premezclas y harinas', 'A0118401': 'Alfajores', 'A0118402': 'Cereales, copos y barras',
    'A0111207': 'Cereales, copos y barras', 'A0111205': 'Cereales, copos y barras',
    'A0111311': 'Tapas, masas y pizzas', 'A0111312': 'Tapas, masas y pizzas', 'A0111307': 'Tapas, masas y pizzas',
    'A0111308': 'Tapas, masas y pizzas', 'A0111303': 'Pastas frescas y ñoquis', 'A0111309': 'Pastas frescas y ñoquis',
    'A0111315': 'Pastas frescas y ñoquis', 'A0213301': 'Cerveza'}
with zipfile.ZipFile(F_GAS) as z:
    g = pd.read_csv(z.open('engho2018_gastos.txt'), sep='|', encoding='latin-1', dtype={'articulo': str},
                    usecols=['articulo', 'division', 'monto', 'pondera', 'region'])
with zipfile.ZipFile(F_ART) as z:
    art = pd.read_csv(z.open('engho2018_articulos.txt'), sep='|', encoding='latin-1', dtype=str) \
        .set_index('articulo')['articulo_desc']
g['gasto'] = g['monto'] * g['pondera']
tot_ab = g.loc[g['division'].isin(['A01', 'A02']), 'gasto'].sum()
s = g[g['articulo'].isin(TIPO_ENGHO)].copy(); s['tipo'] = s['articulo'].map(TIPO_ENGHO)
REG = {1: 'GBA', 2: 'Pampeana', 3: 'Noroeste', 4: 'Noreste', 5: 'Cuyo', 6: 'Patagonia'}
pt = s.groupby('tipo')['gasto'].sum()
PON = pd.DataFrame({'pct_canasta_TACC': pt / pt.sum() * 100, 'pct_alimentos_y_bebidas': pt / tot_ab * 100})
for r, n in REG.items():
    x = s[s['region'] == r].groupby('tipo')['gasto'].sum()
    PON[f'pct_canasta_{n}'] = x / x.sum() * 100
PON = PON.sort_values('pct_canasta_TACC', ascending=False).round(2).reset_index()
PA = s.groupby(['tipo', 'articulo'])['gasto'].sum().reset_index()
PA['articulo_desc'] = PA['articulo'].map(art).str.encode('latin-1', 'ignore').str.decode('utf-8', 'ignore')
PA['pct_canasta_TACC'] = (PA['gasto'] / pt.sum() * 100).round(2)
PA = PA.drop(columns='gasto').sort_values('pct_canasta_TACC', ascending=False)

# ── Serie mensual de cobertura (sucursales por día, promedio del mes) de los pares y los sin TACC multicadena ────────
ids_ser = list(dict.fromkeys([i for p in PARES for i in (p[2], p[3])] +
                             list(H[(H['lado'] == 'sin TACC') & (H['sucursales'] >= 300) & (H['cadenas'] >= 5)]
                                  .sort_values(['tipo', 'sucursales'], ascending=[True, False])['id_producto'])))
SS = SER.reindex(ids_ser).round(0)
SS.insert(0, 'descripcion', Hi.reindex(ids_ser)['descripcion'].values)
SS.insert(0, 'lado', Hi.reindex(ids_ser)['lado'].values)
SS.insert(0, 'tipo', Hi.reindex(ids_ser)['tipo'].values)
SS.insert(0, 'en_par_recomendado', [i in {x for p in PARES for x in (p[2], p[3])} for i in ids_ser])

MULTI = H[(H['lado'] == 'sin TACC') & (H['sucursales'] >= 300) & (H['cadenas'] >= 5)].copy()
MULTI['tiene_par_misma_marca'] = MULTI['id_producto'].isin(set(PR['sin_id']))
colsH = ['tipo', 'forma', 'lado', 'id_producto', 'descripcion', 'marca', 'gramos', 'sucursales', 'cadenas', 'provincias',
         'precio_100g', 'primer_mes', 'meses_con_dato', 'suc_mediana_mes', 'hueco_dic23_abr24', 'cob_2022_sobre_2026',
         'suc_min_2022', 'mes_min_2022', 'suc_min_2024_05', 'min_sobre_mediana_2024_05', 'mes_min_2024_05',
         'estable_desde_2022', 'estable_desde_2024_05', 'nb09_comercios', 'nb09_provincias', 'nb09_pct_cadena_principal',
         'cadenas_detalle']
LEEME = pd.DataFrame({'hoja': [
    'Pares_recomendados', 'Sin_TACC_multicadena', 'Pares_misma_marca_auto', 'Historia_candidatos', 'Cobertura_mensual',
    'Ponderaciones_ENGHo', 'Ponderaciones_ENGHo_articulos', 'Referencias_con_TACC', 'Todos_por_tipo', '', 'Fuentes', '',
    '', 'Reglas', '', '', '', 'Cuidado'], 'contenido': [
    'Pares de la MISMA marca y producto (con TACC / sin TACC) revisados a mano: cobertura el 22/10/2026, brecha ese día, '
    'desde cuándo existe cada lado y si su cobertura es estable desde 2024-05',
    'Productos sin TACC con >= 300 sucursales y >= 5 cadenas el 22/10/2026, con su historia 2022-2026',
    'Emparejamiento automático por marca, tipo, forma (fideos) y tamaño (0,5-2x): punto de partida, revisar a mano',
    'Todos los candidatos (sin TACC y con TACC a base de trigo, >= 50 sucursales el 22/10) con su historia 2022-2026',
    'Sucursales con precio por día, en promedio, en cada mes 2022-01 -> 2026-10 (panel mensual del nb09)',
    'Gasto de los hogares por tipo (ENGHo 2017/18, hogar promedio, total y por región): ponderaciones posibles',
    'Ídem por artículo de la ENGHo',
    'Productos con TACC (trigo, marca no celíaca) con más cobertura por tipo',
    'Catálogo completo del 22/10/2026 clasificado por tipo y lado', '',
    'SEPA diario del 22/10/2026 (16 comercios, 2.954 sucursales; no trae los comercios 19, 1000, 2000, 2002, 2005, '
    '2007, 2013, 3001 ni Libertad) para nombres, presentaciones, cobertura y precios',
    'Panel mensual y ranking del nb09 (trazabilidad_productos_20220101_20261022) para la historia (ago-oct 2026 = SEPA '
    'diario, sin 8 comercios)',
    'ENGHo 2017/18 (INDEC, microdatos de gastos) para las ponderaciones',
    'sin TACC = la descripción del SEPA o del Maestro dice sin TACC / sin gluten / libre de gluten (variantes S TACC, '
    'S/TACC, S-TACC, LIBRE D/GLUTEN, Glutal...); "C/GLUTEN" y galletitas DE gluten son con TACC',
    'con TACC = sin esa marca, sin ingredientes naturalmente sin gluten (arroz, maíz, legumbres...) y de una marca que '
    'no sea celíaca',
    'estable_desde_2024_05 = precio todos los meses 2024-05 -> 2026-07, >= 50 sucursales/día y ningún mes por debajo '
    'del 35% de su mediana (estable_desde_2022: 2022-01 -> 2026-07, umbral 25%)',
    'hueco_dic23_abr24 = cobertura mínima de dic-2023 a abr-2024 sobre la mediana del producto (Molinos cambió los '
    'EANs de sus fideos en sep-oct 2023: el producto viejo y el nuevo conviven unos meses)',
    'El Maestro interno puede tener descripciones viejas para EANs reutilizados (7790070336002 = spaghetti sin TACC '
    'Matarazzo según todas las cadenas; el Maestro dice "Mostachol 5 Cereales"). Usar la descripción del SEPA.']})

def formato(ws):
    fill = PatternFill('solid', start_color='375623', end_color='375623')
    for c in ws[1]:
        c.fill = fill; c.font = Font(bold=True, color='FFFFFF'); c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'
    for i in range(1, ws.max_column + 1):
        h = str(ws.cell(1, i).value or '')
        ws.column_dimensions[get_column_letter(i)].width = 46 if any(k in h for k in ('descripcion', 'nota', 'detalle',
                                                                                      'contenido', 'desc')) else 14

with pd.ExcelWriter(F_OUT, engine='openpyxl') as w:
    LEEME.to_excel(w, sheet_name='Leeme', index=False)
    REC.to_excel(w, sheet_name='Pares_recomendados', index=False)
    MULTI.sort_values(['tipo', 'sucursales'], ascending=[True, False])[colsH + ['tiene_par_misma_marca']] \
        .to_excel(w, sheet_name='Sin_TACC_multicadena', index=False)
    PR.drop(columns=[c for c in ['suc_min'] if c in PR]).to_excel(w, sheet_name='Pares_misma_marca_auto', index=False)
    H.sort_values(['tipo', 'lado', 'sucursales'], ascending=[True, False, False])[colsH] \
        .to_excel(w, sheet_name='Historia_candidatos', index=False)
    SS.reset_index().to_excel(w, sheet_name='Cobertura_mensual', index=False)
    PON.to_excel(w, sheet_name='Ponderaciones_ENGHo', index=False)
    PA.to_excel(w, sheet_name='Ponderaciones_ENGHo_articulos', index=False)
    REF.to_excel(w, sheet_name='Referencias_con_TACC', index=False)
    TODOS.to_excel(w, sheet_name='Todos_por_tipo', index=False)
    for ws in w.sheets.values():
        formato(ws)
    w.sheets['Leeme'].column_dimensions['B'].width = 140
print(REC[['tipo', 'par', 'sin_TACC_sucursales_22oct', 'sin_TACC_cadenas_22oct', 'con_TACC_sucursales_22oct',
           'comparacion', 'precio_sin_TACC', 'precio_con_TACC', 'brecha_22oct_%', 'sin_TACC_desde', 'con_TACC_desde',
           'sin_TACC_estable_desde_2024_05', 'con_TACC_estable_desde_2024_05']].to_string(index=False))
print(); print(PON.to_string(index=False))
print(f'\nGuardado: {F_OUT}')
