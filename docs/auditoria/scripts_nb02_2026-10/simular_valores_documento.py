# Corre el bloque AMBA de la CELDA 7, la CELDA 18 (para det_barrio) y la CELDA 20 REALES del nb02 generado sobre
# los datos de una corrida, reconstruidos desde su canasta_analisis_YYYY-MM.xlsx (Sucs_*, Prov_*, Evolucion_IPC).
# Sirve para ver la hoja Valores_Documento que daria el codigo actual sin re-correr el notebook en Colab.
# Trabaja sobre una COPIA del Excel en <carpeta de trabajo>; el original no se toca.
# Uso: python -I docs/auditoria/scripts_nb02_2026-10/simular_valores_documento.py <repo> <canasta_analisis.xlsx>
#        <carpeta de trabajo> [PERIODO, default 2026-10] [mes parcial: 1/0, default 1]
# Resultado del 2026-10-16 (oct-26 al 15/10): AMBA de la Media 1.405 sucursales (sin Toledo, Coop. Obrera ni
# La Anonima); acumulados ene-24 -> sep-26: Media +162,77% vs IPC Alimentos +166,35%; P5-P95 6,5%.
import contextlib, io, json, pathlib, shutil, sys

import numpy as np
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')
repo, xls_in, out_dir = pathlib.Path(sys.argv[1]), sys.argv[2], pathlib.Path(sys.argv[3])
PERIODO = sys.argv[4] if len(sys.argv) > 4 else '2026-10'
PARCIAL = (sys.argv[5] != '0') if len(sys.argv) > 5 else True

nb = json.load(open(repo / 'notebooks/02_evolucion_canasta_representativa.ipynb', encoding='utf-8'))
CELDAS = [''.join(c['source']) for c in nb['cells'] if c['cell_type'] == 'code']
celda = lambda t: next(c for c in CELDAS if t in c[:300])
bloque = lambda s, a, b: s[s.index(a):s.index(b, s.index(a))]
C3, C6, C7, C18, C20 = (celda('# CELDA 3 —'), celda('# CELDA 6 —'), celda('# CELDA 7 —'),
                        celda('# CELDA 18 —'), celda('# CELDA 20 — Valores'))

env = dict(np=np, pd=pd)
exec(bloque(C3, 'CANASTA_NAMES = {', 'CANASTA_LINESTYLES'), env)
exec(bloque(C6, 'def _pmean(_s):', 'ref_e ='), env)
SH = env['CANASTA_SHORT']

out_dir.mkdir(parents=True, exist_ok=True)
xls = out_dir / pathlib.Path(xls_in).name
shutil.copy(xls_in, xls)
wb = openpyxl.load_workbook(xls)
if 'Valores_Documento' in wb.sheetnames:
    del wb['Valores_Documento']
wb.save(xls)
hojas = wb.sheetnames
ACT = [c for c in env['CANASTA_NAMES'] if f'Sucs_{SH[c]}' in hojas]

geo, prov, pais, pais_p = {}, {}, {}, {}
for c in ACT:
    s = pd.read_excel(xls, sheet_name=f'Sucs_{SH[c]}')
    e = dict(env, suc_geo_clean=s)
    with contextlib.redirect_stdout(io.StringIO()):
        exec(bloque(C7, '# ── AMBA por coordenadas', '# ── Calcular canasta por sucursal'), e)
    geo[c] = e['suc_geo_clean']
    p = pd.read_excel(xls, sheet_name=f'Prov_{SH[c]}')
    prov[c] = p
    # el pais del Cuadro 1 se recupera de cualquier provincia: valor / (1 + vs_promedio)
    pais[c] = float((p['canasta_total'] / (1 + p['vs_promedio_%'] / 100)).median())
    pais_p[c] = float((p['canasta_total_prom'] / (1 + p['vs_promedio_prom_%'] / 100)).median())
evo = pd.read_excel(xls, sheet_name='Evolucion_IPC')
evo['mes'] = evo['mes'].astype(str)
serie = {}
for c in ACT:
    n = SH[c]
    serie[c] = (evo[['mes', f'canasta_{n}', f'var_{n}_%', f'canasta_{n}_prom', f'var_{n}_prom_%']]
                .dropna(subset=[f'canasta_{n}'])
                .rename(columns={f'canasta_{n}': 'canasta_nacional_ponderada', f'var_{n}_%': 'variacion_mensual_%',
                                 f'canasta_{n}_prom': 'canasta_nacional_ponderada_prom',
                                 f'var_{n}_prom_%': 'variacion_mensual_prom_%'}).reset_index(drop=True))
ipc = (evo[['mes', 'ipc_general', 'ipc_general_var_%', 'ipc_alimentos', 'ipc_alimentos_var_%']]
       .dropna(subset=['ipc_general']).reset_index(drop=True))
# productos de la Media: la hoja Productos de la 19b; si no esta, el maximo de productos propios
_n_media = (len(pd.read_excel(xls, sheet_name='Productos')) if 'Productos' in hojas
            else int(geo['cantidad_03']['productos_propios'].max()))
CANASTAS = {c: {str(k): None for k in range(_n_media if c == 'cantidad_03' else 1)} for c in ACT}

env.update(canasta_geo_dict=geo, prom_nac_dict=pais, prom_nac_prom_dict=pais_p, serie_prov_dict=prov,
           serie_nac_dict=serie, ipc=ipc, CANASTAS_ACTIVAS=ACT, CANASTAS=CANASTAS,
           MIN_SUCURSALES_RANKING=10, OUTPUT_DIR=out_dir, MES=PERIODO[5:] + PERIODO[:4], PERIODO=PERIODO,
           MES_INICIO_GRAFICO='2024-01', NOMBRE_MES_TITLE=PERIODO,
           SUFIJO_PARCIAL=' (parcial)' if PARCIAL else '', MES_PARCIAL=PARCIAL)
print(f"pais (Cuadro 1) Media: {pais['cantidad_03']:,.0f} | AMBA Media: {int(geo['cantidad_03']['amba'].sum()):,} sucursales")
with contextlib.redirect_stdout(io.StringIO()):
    exec(C18, env)
exec(C20, env)
