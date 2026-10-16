# -*- coding: utf-8 -*-
"""Test de los numeros del documento tecnico del nb02: AMBA, referencias, acumulados y formato.

Corre las celdas REALES del notebook generado (`02_evolucion_canasta_representativa.ipynb`): el
bloque AMBA de la CELDA 7, la CELDA 16 (rankings de cadenas), la CELDA 18 (barrios CABA) y la
CELDA 20 (Valores_Documento), sobre sucursales sinteticas puestas donde muerden: Toledo en Mar del
Plata, la Cooperativa Obrera en Bahia Blanca, DIA en el conurbano y en Pergamino, Vea en Mendoza,
una sucursal atipica 20% abajo, y la serie con un mes (oct-26, parcial) que el IPC no publico.

Comprueba los defectos de la corrida de oct-26 (2026-10-16):
  1. AMBA = CABA + Buenos Aires a <= 90 km del Obelisco (La Plata, Campana y Zarate adentro; Mar del
     Plata, Lobos y Pergamino afuera). Antes era CABA + TODA la provincia: Toledo y la Cooperativa
     Obrera salian en el ranking "AMBA".
  2. Acumulados de las canastas y del IPC en el MISMO periodo (hasta el ultimo mes con IPC); el mes
     sin IPC va aparte. Antes: canasta a oct-26 contra IPC a sep-26.
  3. Una sola referencia del pais (Cuadro 1) en las CELDAS 16, 18 y 20; cadenas y barrios por
     mediana de sus sucursales; CABA = su fila del Cuadro 1. Antes habia cuatro "pais" distintos y
     la CELDA 20 usaba promedios por cadena y barrio.
  4. Tantas caras como baratas en el resumen (antes 3 caras y 1 barata).
  5. Dispersion entre sucursales P5-P95 (el min-max queda marcado "no publicar" con la sucursal).
  6. Porcentajes con coma decimal y signo correcto (antes "+8.7%" con abs() en la prima celiaca).

    python notebooks/test_valores_documento.py     -> RESULTADO: OK / FALLA
"""
import contextlib, io, json, pathlib, re, tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
import openpyxl
import pandas as pd

NB = pathlib.Path(__file__).with_name('02_evolucion_canasta_representativa.ipynb')
CELDAS = [''.join(c['source']) for c in json.load(open(NB, encoding='utf-8'))['cells']
          if c['cell_type'] == 'code']


def celda(tag):
    return next(c for c in CELDAS if tag in c[:300])


def bloque(src, ini, fin):
    i0 = src.index(ini)
    return src[i0:src.index(fin, i0)]


C3, C6, C7 = celda('# CELDA 3 —'), celda('# CELDA 6 —'), celda('# CELDA 7 —')
C16, C18, C20 = celda('# CELDA 16 —'), celda('# CELDA 18 —'), celda('# CELDA 20 — Valores')

fallas = []


def chequear(ok, msg):
    print(('  OK    ' if ok else '  FALLA ') + msg)
    if not ok:
        fallas.append(msg)


def correr(src, env, nombre):
    _out = io.StringIO()
    with contextlib.redirect_stdout(_out):
        exec(compile(src, nombre, 'exec'), env)
    return _out.getvalue()


ACTIVAS = ['cantidad_01', 'cantidad_02', 'cantidad_03', 'cantidad_04', 'cantidad_05', 'cantidad_06']
BASE = dict(np=np, pd=pd)
exec(bloque(C3, 'CANASTA_NAMES = {', 'CANASTA_LINESTYLES'), BASE)
exec(bloque(C6, 'def _pmean(_s):', 'ref_e ='), BASE)

# ── Sucursales sinteticas: (cadena, provincia, lat, lon, n, nivel) ───────────
PUNTOS = [
    ('DIA', 'CABA', -34.580, -58.420, 4, 1000), ('DIA', 'CABA', -34.615, -58.445, 4, 1004),
    ('DIA', 'CABA', -34.630, -58.465, 4, 998),
    ('DIA', 'Buenos Aires', -34.700, -58.580, 8, 995),          # La Matanza
    ('DIA', 'Buenos Aires', -33.894, -60.569, 10, 990),         # Pergamino (fuera del AMBA)
    ('Jumbo', 'CABA', -34.560, -58.450, 6, 1040), ('Jumbo', 'CABA', -34.582, -58.425, 6, 1036),
    ('Coto', 'CABA', -34.615, -58.445, 6, 1010), ('Coto', 'CABA', -34.630, -58.465, 6, 1012),
    ('Coto', 'Buenos Aires', -34.720, -58.250, 8, 1006),        # Quilmes
    ('Toledo', 'Buenos Aires', -38.000, -57.550, 12, 1020),     # Mar del Plata
    ('Cooperativa Obrera', 'Buenos Aires', -38.720, -62.270, 12, 1015),   # Bahia Blanca
    ('Vea', 'Mendoza', -32.890, -68.840, 12, 975),
]
filas = []
rng = np.random.default_rng(7)
for cad, prov, lat, lon, n, nivel in PUNTOS:
    for k in range(n):
        filas.append({'id_comercio': cad[:3], 'id_bandera': '1', 'id_sucursal': f'{lat}{k}',
                      'cadena': cad, 'PROVINCIA_NORM': prov,
                      'sucursales_nombre': f'{cad} {k}', 'sucursales_latitud': lat + rng.normal(0, 0.002),
                      'sucursales_longitud': lon + rng.normal(0, 0.002),
                      'canasta_total': nivel * (1 + rng.normal(0, 0.004)) * 680,
                      'productos_propios': 69})
suc = pd.DataFrame(filas)
suc.loc[len(suc)] = {'id_comercio': 'Hip', 'id_bandera': '4', 'id_sucursal': 'x', 'cadena': 'Hiper ChangoMas',
                     'PROVINCIA_NORM': 'Buenos Aires', 'sucursales_nombre': 'San Martin',
                     'sucursales_latitud': -34.574, 'sucursales_longitud': -58.530,
                     'canasta_total': 548469.0, 'productos_propios': 66}          # atipica, 20% abajo
suc['canasta_total_prom'] = suc['canasta_total'] * 0.998

# ── 1. Bloque AMBA real de la CELDA 7 ────────────────────────────────────────
print('1. AMBA por coordenadas (CELDA 7)')
casos = pd.DataFrame([
    ('CABA', -34.60, -58.40, True), ('Buenos Aires', -34.92, -57.95, True),     # La Plata
    ('Buenos Aires', -34.16, -58.96, True), ('Buenos Aires', -34.10, -59.03, True),   # Campana, Zarate
    ('Buenos Aires', -38.00, -57.55, False), ('Buenos Aires', -35.18, -59.10, False),  # MdP, Lobos
    ('Buenos Aires', -33.89, -60.57, False), ('Santa Fe', -34.60, -58.40, False),
], columns=['PROVINCIA_NORM', 'sucursales_latitud', 'sucursales_longitud', 'esperado'])
env7 = dict(BASE, suc_geo_clean=casos.copy())
correr(bloque(C7, '# ── AMBA por coordenadas', '# ── Calcular canasta por sucursal'), env7, '<CELDA 7 AMBA>')
_res = env7['suc_geo_clean']
chequear(_res['amba'].tolist() == casos['esperado'].tolist(),
         f"CABA, La Plata, Campana y Zarate adentro; Mar del Plata, Lobos, Pergamino y otra provincia afuera "
         f"({_res['amba'].tolist()})")
env7['suc_geo_clean'] = suc.copy()
correr(bloque(C7, '# ── AMBA por coordenadas', '# ── Calcular canasta por sucursal'), env7, '<CELDA 7 AMBA>')
suc = env7['suc_geo_clean']
chequear(not suc.loc[suc['cadena'].isin(['Toledo', 'Cooperativa Obrera']), 'amba'].any(),
         'Toledo (Mar del Plata) y la Cooperativa Obrera (Bahia Blanca) quedan fuera del AMBA')

# Media = las sucursales; las demas canastas, escaladas (solo la Media entra en la geografia)
ESC = {'cantidad_01': 0.42, 'cantidad_02': 0.74, 'cantidad_03': 1.0, 'cantidad_04': 1.44,
       'cantidad_05': 0.95, 'cantidad_06': 0.69}       # celiaca por DEBAJO de la Media: signo negativo
geo = {}
for c, f in ESC.items():
    g = suc.copy()
    g['canasta_total'] *= f
    g['canasta_total_prom'] *= f
    geo[c] = g
PAIS = {c: 690000.0 * f for c, f in ESC.items()}      # Cuadro 1: distinto de la mediana de sucursales
PAIS_PROM = {c: 689000.0 * f for c, f in ESC.items()}
serie_prov = {c: (geo[c].groupby('PROVINCIA_NORM')
                  .agg(canasta_total=('canasta_total', 'median'),
                       canasta_total_prom=('canasta_total_prom', BASE['_pmean']))
                  .reset_index().rename(columns={'PROVINCIA_NORM': 'provincia'})) for c in ACTIVAS}

meses = pd.period_range('2024-01', '2026-10', freq='M').astype(str)
serie_nac = {}
for i, c in enumerate(ACTIVAS):
    v = 250000 * ESC[c] * np.cumprod(np.r_[1.0, np.full(len(meses) - 1, 1.03 + 0.001 * i)])
    d = pd.DataFrame({'mes': meses, 'canasta_nacional_ponderada': v, 'canasta_nacional_ponderada_prom': v * 0.99})
    d['variacion_mensual_%'] = d['canasta_nacional_ponderada'].pct_change() * 100
    d['variacion_mensual_prom_%'] = d['canasta_nacional_ponderada_prom'].pct_change() * 100
    serie_nac[c] = d
mi = pd.period_range('2017-01', '2026-09', freq='M').astype(str)
ipc = pd.DataFrame({'mes': mi, 'ipc_general': 100 * 1.035 ** np.arange(len(mi)),
                    'ipc_alimentos': 100 * 1.033 ** np.arange(len(mi))})
ipc['ipc_general_var_%'] = ipc['ipc_general'].pct_change() * 100
ipc['ipc_alimentos_var_%'] = ipc['ipc_alimentos'].pct_change() * 100

OUT = pathlib.Path(tempfile.mkdtemp(prefix='test_valores_doc_'))
COMUN = dict(BASE, canasta_geo_dict=geo, prom_nac_dict=PAIS, prom_nac_prom_dict=PAIS_PROM,
             MIN_SUCURSALES_RANKING=10, OUTPUT_DIR=OUT, MES='102026', PERIODO='2026-10',
             NOMBRE_MES_TITLE='Octubre 2026', SUFIJO_PARCIAL=' (parcial: 15/31 días)', MES_PARCIAL=True)

# ── 2. CELDA 16: rankings de cadenas ─────────────────────────────────────────
print('2. CELDA 16 (rankings de cadenas)')
_savefig, _show = plt.savefig, plt.show
plt.savefig = lambda *a, **k: None
plt.show = lambda *a, **k: None
try:
    env16 = dict(COMUN, plt=plt, mticker=mticker, CANASTAS_ACTIVAS=['cantidad_03'])
    correr(C16, env16, '<CELDA 16>')
    figs = {f.axes[0].get_title(): f.axes[0] for f in map(plt.figure, plt.get_fignums())}
finally:
    plt.savefig, plt.show = _savefig, _show
ax_n, ax_a = figs['Ranking nacional [Media/mediana]'], figs['Ranking AMBA [Media/mediana]']
chequear(abs(ax_n.lines[0].get_xdata()[0] - PAIS['cantidad_03']) < 1e-6
         and ax_n.get_legend().get_texts()[0].get_text().startswith('Nacional (Cuadro 1): $690.000'),
         'ranking nacional: la linea es el pais del Cuadro 1 ($690.000), no la mediana de sucursales')
_cad_amba = {t.get_text().split('  (')[0] for t in ax_a.get_yticklabels()}
_med_amba = suc.loc[suc['amba'], 'canasta_total'].median()
chequear(_cad_amba == {'DIA', 'Coto', 'Jumbo'},
         f'ranking AMBA: DIA, Coto y Jumbo, sin Toledo ni la Cooperativa Obrera ({sorted(_cad_amba)})')
chequear(abs(ax_a.lines[0].get_xdata()[0] - _med_amba) < 1e-6
         and ax_a.get_legend().get_texts()[0].get_text().startswith('AMBA: $'),
         'ranking AMBA: la linea es la mediana del AMBA y dice "AMBA" (antes decia "Nacional")')
plt.close('all')

# ── 3. CELDA 18: barrios CABA ────────────────────────────────────────────────
print('3. CELDA 18 (barrios CABA)')
env18 = dict(COMUN, CANASTAS_ACTIVAS=['cantidad_03'])
out18 = correr(C18, env18, '<CELDA 18>')
chequear('Pais (Cuadro 1, ponderado): $690.000' in out18 and 'Pais (Cuadro 1, ponderado): $689.000' in out18,
         'barrios: el pais es el del Cuadro 1 en la mediana y en el promedio')
_caba = suc[suc['PROVINCIA_NORM'] == 'CABA']
chequear(f"CABA (Cuadro 1): {env18['fmtb'](_caba['canasta_total'].median())}" in out18,
         'barrios: CABA = mediana de sus sucursales (su fila del Cuadro 1)')

# ── 4. CELDA 20: Valores_Documento ───────────────────────────────────────────
print('4. CELDA 20 (Valores_Documento)')
XLS = OUT / 'canasta_analisis_2026-10.xlsx'
_wb = openpyxl.Workbook(); _wb.active.title = 'analisis_detallado'; _wb.save(XLS)
CANASTAS = {c: {str(k): ('p', 1, 'x') for k in range(69)} for c in ACTIVAS}
env20 = dict(COMUN, CANASTAS_ACTIVAS=ACTIVAS, CANASTAS=CANASTAS, serie_nac_dict=serie_nac,
             serie_prov_dict=serie_prov, ipc=ipc, MES_INICIO_GRAFICO='2024-01',
             det_barrio=env18['det_barrio'])
out20 = correr(C20, env20, '<CELDA 20>')
correr(C20, env20, '<CELDA 20 otra vez>')              # re-correr reemplaza la hoja
wb = openpyxl.load_workbook(XLS)
chequear(wb.sheetnames == ['analisis_detallado', 'Valores_Documento'],
         f'la hoja se agrega una sola vez aunque se re-corra la celda ({wb.sheetnames})')
vd = pd.read_excel(XLS, sheet_name='Valores_Documento')
fila = lambda sec, var: vd[(vd['Seccion'] == sec) & (vd['Variable'] == var)]
num = lambda sec, var: float(fila(sec, var)['Valor_numero'].iloc[0]) if len(fila(sec, var)) else float('nan')
txt = lambda sec, var: str(fila(sec, var)['Valor_LaTeX'].iloc[0]) if len(fila(sec, var)) else ''

# 4a. AMBA
_amba_rows = set(vd.loc[vd['Seccion'] == 'Cadenas_AMBA', 'Variable'])
chequear({'DIA', 'Coto', 'Jumbo'} <= _amba_rows and not {'Toledo', 'Cooperativa Obrera'} & _amba_rows,
         'Cadenas_AMBA: DIA, Coto y Jumbo; sin Toledo ni la Cooperativa Obrera')
_dia_amba = suc[suc['amba'] & (suc['cadena'] == 'DIA')]['canasta_total'].median()
chequear(abs(num('Cadenas_AMBA', 'DIA') - round(_dia_amba)) < 1
         and abs(num('Cadenas_AMBA', 'DIA — vs AMBA') - round((_dia_amba / _med_amba - 1) * 100, 2)) < 1e-9,
         'Cadenas_AMBA: DIA = mediana de sus sucursales del AMBA, vs la mediana del AMBA')

# 4b. Referencia unica del pais y medianas
_dia = suc[suc['cadena'] == 'DIA']['canasta_total'].median()
chequear(abs(num('Cadenas_Nacional', 'DIA') - round(_dia)) < 1
         and abs(num('Cadenas_Nacional', 'DIA — vs país') - round((_dia / PAIS['cantidad_03'] - 1) * 100, 2)) < 1e-9,
         'Cadenas_Nacional: mediana por cadena, vs el pais del Cuadro 1')
chequear(abs(num('Barrios_CABA', 'CABA (Cuadro 1)') - round(_caba['canasta_total'].median())) < 1
         and abs(num('Provincias', 'País (Cuadro 1)') - PAIS['cantidad_03']) < 1,
         'Barrios_CABA: CABA = su fila del Cuadro 1; Provincias trae el pais del Cuadro 1')
_bar = _caba.assign(b=[env18['det_barrio'](a, o) for a, o in
                       zip(_caba['sucursales_latitud'], _caba['sucursales_longitud'])])
_b1 = _bar[_bar['b'] != 'Sin clasificar'].groupby('b')['canasta_total'].median()
_bv = vd[(vd['Seccion'] == 'Barrios_CABA') & vd['Variable'].isin(_b1.index)].set_index('Variable')['Valor_numero']
chequear(len(_bv) >= 3 and all(abs(float(_bv[b]) - round(_b1[b])) < 1 for b in _bv.index),
         f'Barrios_CABA: mediana por barrio ({len(_bv)} barrios)')

# 4c. Acumulados en el mismo periodo
_s = serie_nac['cantidad_03'].set_index('mes')['canasta_nacional_ponderada']
_i = ipc.set_index('mes')
chequear(abs(num('Acumulados', 'Media 2024-01 a 2026-09') - round((_s['2026-09'] / _s['2024-01'] - 1) * 100, 2)) < 1e-9
         and abs(num('Acumulados', 'IPC Nivel General 2024-01 a 2026-09')
                 - round((_i.loc['2026-09', 'ipc_general'] / _i.loc['2024-01', 'ipc_general'] - 1) * 100, 2)) < 1e-9
         and abs(num('Acumulados', 'IPC Alimentos 2024-01 a 2026-09')
                 - round((_i.loc['2026-09', 'ipc_alimentos'] / _i.loc['2024-01', 'ipc_alimentos'] - 1) * 100, 2)) < 1e-9,
         'acumulados de la Media y del IPC hasta 2026-09 (el ultimo mes con IPC)')
chequear(abs(num('Acumulados', 'Media 2024-01 a 2026-10 (sin IPC; no comparar)')
             - round((_s['2026-10'] / _s['2024-01'] - 1) * 100, 2)) < 1e-9
         and 'Item 6: acumulados 2024-01 -> 2026-09' in out20,
         'el mes sin IPC (oct-26) va aparte y marcado')

# 4d. Caras y baratas simetricas
_it4 = out20[out20.index('Item 4'):out20.index('Item 5')]
_nac4, _, _amba4 = _it4.partition('\n  AMBA')
chequear(_nac4.count('CARA ') == _nac4.count('BARATA ') == 3 and _amba4.count('CARA ') == _amba4.count('BARATA ') == 1,
         'resumen: 3 caras y 3 baratas a nivel nacional; 1 y 1 en el AMBA (3 cadenas)')

# 4e. Dispersion P5-P95
_v = suc['canasta_total']
_p5, _p95 = _v.quantile(0.05), _v.quantile(0.95)
chequear(abs(num('Provincias', 'Dispersión inter-sucursal P5-P95 (%)') - round((_p95 / _p5 - 1) * 100, 1)) < 1e-9
         and not fila('Provincias', 'Dispersión inter-sucursal (%)').shape[0]
         and fila('Provincias', 'Dispersión inter-sucursal mín-máx (%) — no publicar').shape[0] == 1,
         'dispersion entre sucursales P5-P95; el min-max queda marcado "no publicar"')
chequear('Hiper ChangoMas San Martin' in txt('Provincias', 'Sucursal más barata (cuál)'),
         'la sucursal minima se identifica (cadena, nombre, productos propios)')

# 4f. Formato
_pcts = [s for s in vd['Valor_LaTeX'].astype(str) if s.endswith('%')]
_mal = [s for s in _pcts if not re.fullmatch(r'[+-]?\d+(,\d+)?%', s)]
chequear(len(_pcts) > 20 and not _mal, f'todos los porcentajes con coma decimal ({len(_pcts)}; mal: {_mal[:5]})')
chequear(txt('Especiales', 'Prima celíaca (%)') == '-5,00%',
         f"prima celiaca con su signo (celiaca 5% abajo de la Media -> -5,00%; "
         f"da {txt('Especiales', 'Prima celíaca (%)')})")
chequear('Brecha relativa:       242,9%' in out20 and 'Prima celíaca:   -5,00%' in out20,
         'el resumen impreso usa el mismo formato')

print('\nRESULTADO:', 'OK' if not fallas else f'FALLA ({len(fallas)})')
raise SystemExit(1 if fallas else 0)
