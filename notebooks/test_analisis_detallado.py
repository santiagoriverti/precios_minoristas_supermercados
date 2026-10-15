# -*- coding: utf-8 -*-
"""Test de la CELDA 19c del nb02: hoja 'analisis_detallado' (indices por rubro, region, provincia,
cadena y tipo de sucursal, mensuales y semanales).

Corre las celdas REALES del notebook generado (`02_evolucion_canasta_representativa.ipynb`): CELDA 5,
CELDA 9, la 19b y la 19c (mas los nombres de la CELDA 3, `asignar_cadena` de la CELDA 4, `_pmean` de
la CELDA 6 y el formato de la CELDA 19) sobre ZIPs semestrales sinteticos: mar-26 a jun-26 completos
en 2026A.zip y jul-26 al 15 en 2026B.zip. Con DETALLE_MESES = 3 la hoja muestra abr, may y jun
(base mar-26 = 100), jul parcial y las semanas 1S-4S de junio y 1S-2S de julio (la 3S de julio
tiene un solo dia: no va).

Comprueba, contra un calculo independiente hecho con listas y statistics.median, que:
  1. el Total es la serie oficial de la CELDA 9 (indice y variacion, meses y semanas);
  2. cada rubro y categoria es el indice encadenado de las medianas nacionales de sus productos, la
     Pond es su participacion en el costo del mes base y el Total = suma de Pond x indice / 100;
  3. cada region, provincia, cadena y tipo es el indice encadenado de la mediana de las
     observaciones del grupo (con un producto que falta en un grupo un mes: eslabon apareado);
  4. las semanas son indice del mes anterior x canasta de la semana / canasta del mes anterior, y su
     variacion es contra la semana anterior;
  5. los grupos con menos de DETALLE_MIN_SUCURSALES no se muestran y quedan en las notas;
  6. la hoja es la PRIMERA del Excel, la unica pestaña seleccionada, y las demas hojas siguen;
     sus celdas tienen los mismos numeros que `detalle_long`;
  7. la segunda corrida usa el cache (lee del SEPA solo lo que no esta) y da lo mismo; si cambia el
     maestro de sucursales, la clave cambia y relee; con DESAGREGAR_CANASTAS vacio la hoja sale igual;
  8. con el mes en curso COMPLETO (julio al 31) no hay columna parcial, van las 4 semanas de cada mes
     y el cache no guarda el mes en curso.

    python notebooks/test_analisis_detallado.py     -> RESULTADO: OK / FALLA
"""
import calendar, gc, gzip, hashlib, json, pathlib, re, shutil, statistics, tempfile, zipfile
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

NB = pathlib.Path(__file__).with_name('02_evolucion_canasta_representativa.ipynb')
CELDAS = [''.join(c['source']) for c in json.load(open(NB, encoding='utf-8'))['cells']
          if c['cell_type'] == 'code']


def celda(tag):
    return next(c for c in CELDAS if tag in c[:300])


def bloque(src, ini, fin):
    i0 = src.index(ini)
    return src[i0:src.index(fin, i0)]


C3, C4, C5, C6 = celda('# CELDA 3 —'), celda('# CELDA 4 —'), celda('# CELDA 5 —'), celda('# CELDA 6 —')
C9, C19, C19B, C19C = celda('# CELDA 9 —'), celda('# CELDA 19 —'), celda('# CELDA 19b —'), celda('# CELDA 19c —')

fallas = []


def chequear(ok, msg):
    print(('  OK    ' if ok else '  FALLA ') + msg)
    if not ok:
        fallas.append(msg)


def cerca(a, b, tol=1e-7):
    if pd.isna(a) and pd.isna(b):
        return True
    return abs(float(a) - float(b)) <= tol


# ── Datos sinteticos ──────────────────────────────────────────────────────────
MESES = [202603, 202604, 202605, 202606, 202607]
DIAS = {202603: 31, 202604: 30, 202605: 31, 202606: 30, 202607: 15}
PERIODO, MES_ANT, BASE, D_ACT = '2026-07', '2026-06', '2026-03', 15
COLS_M = ['2026-04', '2026-05', '2026-06']
PER_W = ['2026-06-S1', '2026-06-S2', '2026-06-S3', '2026-06-S4', '2026-07-S1', '2026-07-S2']
COMPLETO = False          # el ultimo escenario corre con julio completo (sin columna parcial)
ABR = {3: 'mar', 4: 'abr', 5: 'may', 6: 'jun', 7: 'jul'}


def meses_vis():
    return COLS_M + ([] if COMPLETO else [PERIODO])


def ultimo():
    return PERIODO if COMPLETO else MES_ANT


def etiquetas():
    lab = [f'{ABR[int(m[5:])]}-26' for m in COLS_M] + ([] if COMPLETO else ['jul-26 (al 15/7)'])
    return lab + [f'{w[-1]}S {ABR[int(w[5:7])]}' for w in PER_W]
MIN_SUC = 2
SUCS = [('9', '1', '1', 'AR-C', ('AMBA', 'Supermercado')),
        ('9', '1', '2', 'AR-B', ('AMBA\xa0', 'SUPERMERCADO')),
        ('9', '1', '3', 'AR-X', ('Pampeana', 'Hipermercado')),
        ('15', '1', '1', 'AR-C', ('AMBA', 'Autoservicio')),
        ('15', '1', '2', 'AR-X', ('Pampeana', 'autoservicio')),
        ('15', '1', '3', 'AR-M', ('Cuyo\xa0', 'Autoservicio')),
        ('15', '1', '4', 'AR-M', ('Cuyo', 'Autoservicio')),
        ('12', '1', '1', 'AR-C', ('AMBA', 'Hipermercado')),
        ('12', '1', '2', 'AR-B', ('AMBA', 'Tienda Fisica')),
        ('10', '4', '1', 'AR-S', None),
        ('99', '1', '1', 'AR-B', None),
        ('2', '1', '1', 'AR-R', ('Patagonia', 'Supermercado')),
        ('2', '1', '2', 'AR-U', ('Patagonia', 'Supermercado'))]
CADENA_ESP = {('9', '1'): 'Vea', ('15', '1'): 'DIA', ('12', '1'): 'Coto', ('10', '4'): 'Carrefour (bandera 4)',
              ('99', '1'): 'Cadena 99', ('2', '1'): 'La Anonima'}
PROV_ESP = {'AR-C': 'CABA', 'AR-B': 'Buenos Aires', 'AR-X': 'Córdoba', 'AR-M': 'Mendoza', 'AR-S': 'Santa Fe',
            'AR-R': 'Río Negro', 'AR-U': 'Chubut'}
REGION_FUERA = {'AR-S': 'Pampeana', 'AR-B': 'Buenos Aires s/ubicar'}
TIPO_ESP = {'Supermercado': 'Supermercado', 'SUPERMERCADO': 'Supermercado', 'Hipermercado': 'Hipermercado',
            'Autoservicio': 'Autoservicio', 'autoservicio': 'Autoservicio', 'Tienda Fisica': 'Tienda física'}

# (ean, descripcion, qty, categoria, rubro SEPA)
MEDIA = [('7790001000011', 'Leche Entera 1 Lt', 3, 'Lácteos', 'Frescos'),
         ('7790001000028', 'Yogur Bebible 1 Lt', 2, 'Lácteos', 'Frescos'),
         ('7790001000035', 'Aceite Girasol 900 Ml', 1, 'Aceites', 'Almacén'),
         ('40084107', 'Galletitas de Agua 300 Gr', 2, 'Galletitas', 'Almacén'),
         ('7790001000042', 'Gaseosa Cola 2,25 Lt', 2, 'Gaseosas', 'Bebidas'),
         ('7790001000059', 'Pasta Dental 90 Gr', 1, 'Cuidado Oral', 'Perfumería')]
POPULAR = [('7790001000011', 'Leche Entera 1 Lt', 2, 'Lácteos', 'Frescos'),
           ('7790001000073', 'Arroz Largo Fino 1 Kg', 2, 'Arroz y Legumbres', 'Almacén')]
E_ACEITE = '7790001000035'
SIN_ABRIL = {('12-1-1', E_ACEITE), ('12-1-2', E_ACEITE)}   # Coto sin aceite en abril: eslabon apareado


def sk(s):
    return f'{s[0]}-{s[1]}-{s[2]}'


def generar():
    rng = np.random.default_rng(11)
    eans = sorted({e for e, *_ in MEDIA + POPULAR})
    base = {e: float(rng.integers(600, 4000)) for e in eans}
    obs = []                                     # (sucursal, ean, fecha, precio)
    for si, s in enumerate(SUCS):
        f = rng.uniform(0.85, 1.25)
        for e in eans:
            if rng.random() < 0.1 and e != E_ACEITE:
                continue
            p = round(base[e] * f, 2)
            for ym in MESES:
                for d in range(1, DIAS[ym] + 1):
                    if rng.random() < 0.04:
                        p = round(p * rng.choice([1.02, 1.05, 1.10, 0.97]), 2)
                    if (sk(s), e) in SIN_ABRIL and ym == 202604:
                        continue
                    if rng.random() >= 0.05:
                        obs.append((si, e, ym * 100 + d, p))
    return obs


OBS = generar()


def escribir_zips(sepa_dir):
    serie = defaultdict(dict)
    for si, e, fe, p in OBS:
        serie[(si, e)][fe] = p
    por_zip = defaultdict(dict)
    for ym in MESES:
        zn = '2026A.zip' if ym <= 202606 else '2026B.zip'
        n = DIAS[ym]
        partes = [range(1, min(15, n) + 1)] + ([range(16, n + 1)] if n > 15 else [])
        for k, dias in enumerate(partes, 1):
            fechas = [ym * 100 + d for d in dias]
            lin = [','.join(['id_comercio', 'id_bandera', 'id_sucursal', 'sucursales_provincia', 'id_producto']
                            + [f'precio_{fe}' for fe in fechas])]
            for (si, e), ps in sorted(serie.items()):
                if any(fe in ps for fe in fechas):
                    s = SUCS[si]
                    lin.append(','.join([s[0], s[1], s[2], s[3], e.zfill(13)] +
                                        [('%.2f' % ps[fe]).rstrip('0').rstrip('.') if fe in ps else 'NA'
                                         for fe in fechas]))
            nombre = f'{ym % 100:02d}{ym // 100}_pais_parte{k}COMPLETO.csv.gz'
            por_zip[zn][nombre] = gzip.compress(('\n'.join(lin) + '\n').encode('utf-8'))
    for zn, arch in por_zip.items():
        with zipfile.ZipFile(sepa_dir / zn, 'w', zipfile.ZIP_STORED) as z:
            for fn, b in arch.items():
                z.writestr(fn, b)


def maestro(cambiado=False):
    filas = [{'id_comercio': s[0], 'id_bandera': s[1], 'id_sucursal': s[2], 'REGION': s[4][0],
              'sucursales_tipo': s[4][1]} for s in SUCS if s[4]]
    if cambiado:
        filas[0]['REGION'] = 'Pampeana'
    return pd.DataFrame(filas)


def correr(tmp, nombre, desagregar, cache_dir, contar_lecturas=False, maestro_cambiado=False, sepa='sepa'):
    sepa, out, tdir = tmp / sepa, tmp / f'out_{nombre}', tmp / f'tmp_{nombre}'
    out.mkdir(); tdir.mkdir()
    xls = out / 'canasta_analisis_2026-07.xlsx'
    with pd.ExcelWriter(xls, engine='openpyxl') as w:      # lo que dejo la CELDA 19
        pd.DataFrame({'mes': ['2026-06'], 'x': [1.0]}).to_excel(w, sheet_name='Evolucion_IPC', index=False)
        pd.DataFrame({'x': [1]}).to_excel(w, sheet_name='Serie_precios', index=False)
    env = {'__builtins__': __builtins__, 'np': np, 'pd': pd, 're': re, 'zipfile': zipfile, 'gzip': gzip,
           'shutil': shutil, 'gc': gc, 'hashlib': hashlib, 'Path': Path,
           'SEPA_DIR': sepa, 'OUTPUT_DIR': out, 'CACHE_DIR': cache_dir, 'TMP_DIR': tdir, 'USE_CACHE': True,
           'MES_INICIO_HISTORICO': '2026-03', 'PERIODO': PERIODO, 'ULTIMO_MES': PERIODO,
           'maestro_suc': maestro(maestro_cambiado), 'DESAGREGAR_CANASTAS': desagregar,
           'DETALLE_CANASTA': 'cantidad_03', 'DETALLE_MESES': 3, 'DETALLE_MIN_SUCURSALES': MIN_SUC}
    exec(bloque(C3, 'CANT_COLS =', '# ── Construir CANASTAS'), env)
    env['CANASTAS'] = {'cantidad_03': {e.lstrip('0'): (d[:50], q, c) for e, d, q, c, _ in MEDIA},
                       'cantidad_02': {e.lstrip('0'): (d[:50], q, c) for e, d, q, c, _ in POPULAR}}
    env['RUBRO_EAN'] = {e.lstrip('0'): r for e, _, _, _, r in MEDIA + POPULAR}
    env['CANASTAS_ACTIVAS'] = list(env['CANASTAS'])
    env['CANASTA_EANS_NORM'] = set().union(*[set(c) for c in env['CANASTAS'].values()])
    exec(bloque(C4, 'NOMBRES_COMPUESTOS = {', 'PROV_NORM = {'), env)
    exec(C5, env)
    exec(bloque(C6, 'def _pmean(_s):', 'ref_e ='), env)
    exec(C9, env)
    exec(bloque(C19, 'from openpyxl.styles', 'out_xls ='), env)
    exec(compile(C19B, '<CELDA 19b>', 'exec'), env)
    lecturas = []
    if contar_lecturas:
        orig = env['_leer_mes_desag']
        def contar(lbl, eans=None):
            lecturas.append(lbl)
            return orig(lbl, eans)
        env['_leer_mes_desag'] = contar
    exec(compile(C19C, '<CELDA 19c>', 'exec'), env)
    return env, xls, lecturas


# ── Calculo independiente ─────────────────────────────────────────────────────
def atributos(si, maestro_cambiado=False):
    s = SUCS[si]
    reg = s[4][0].replace('\xa0', '').strip() if s[4] else REGION_FUERA[s[3]]
    if maestro_cambiado and si == 0:
        reg = 'Pampeana'
    return {'region': reg, 'provincia': PROV_ESP[s[3]], 'cadena': CADENA_ESP[(s[0], s[1])],
            'tipo': TIPO_ESP[s[4][1]] if s[4] else 'Sin dato en maestro'}


def periodo_de(fe):
    ym, d = fe // 100, fe % 100
    return f'{ym // 100}-{ym % 100:02d}', f'{ym // 100}-{ym % 100:02d}-S{min((d - 1) // 7, 3) + 1}'


def enlace(a, b, q):
    com = [e for e in q if e in a and e in b]
    if not com:
        return float('nan')
    return sum(q[e] * b[e] for e in com) / sum(q[e] * a[e] for e in com)


def esperado(maestro_cambiado=False):
    q = {e.lstrip('0'): float(qq) for e, _, qq, _, _ in MEDIA}
    meses_g = defaultdict(lambda: defaultdict(list))     # (dim, grupo) -> (ean, mes) -> precios
    sem_g = defaultdict(lambda: defaultdict(list))
    sucs_g = defaultdict(set)                            # (dim, grupo, mes) -> sucursales
    for si, e, fe, p in OBS:
        e = e.lstrip('0')
        if e not in q:
            continue
        mes, sem = periodo_de(fe)
        at = atributos(si, maestro_cambiado)
        for dim, g in [('nacional', 'Total')] + list(at.items()):
            meses_g[(dim, g)][(e, mes)].append(p)
            sem_g[(dim, g)][(e, sem)].append(p)
            sucs_g[(dim, g, mes)].add(si)
    med = lambda dd: {k: statistics.median(v) for k, v in dd.items()}
    per_m = [BASE] + meses_vis()
    res = {}

    def serie(dim, g, eans):
        mm, ww = med(meses_g[(dim, g)]), med(sem_g[(dim, g)])
        qq = {e: q[e] for e in eans}
        col = lambda dd, per: {e: dd[(e, per)] for e in eans if (e, per) in dd}
        idx, var = {}, {}
        idx[BASE] = 100.0
        for k in range(1, len(per_m)):
            l = enlace(col(mm, per_m[k - 1]), col(mm, per_m[k]), qq)
            idx[per_m[k]] = idx[per_m[k - 1]] * l
            var[per_m[k]] = (l - 1) * 100
        r_ant = float('nan')
        for w in PER_W:
            r = enlace(col(mm, MES_ANT), col(ww, w), qq)
            idx[w] = idx[MES_ANT] * r
            var[w] = (r / r_ant - 1) * 100
            r_ant = r
        return idx, var, mm

    eans_all = list(q)
    idx, var, mm = serie('nacional', 'Total', eans_all)
    res[('nacional', 'Total')] = (idx, var)
    t_base = sum(q[e] * mm[(e, BASE)] for e in eans_all)
    pond = {}
    for clave, pos in (('rubro', 4), ('categoria', 3)):
        grupos = defaultdict(list)
        for row in MEDIA:
            grupos[row[pos]].append(row[0].lstrip('0'))
        for g, ee in grupos.items():
            i, v, _ = serie('nacional', 'Total', ee)
            res[(clave, g)] = (i, v)
            pond[(clave, g)] = sum(q[e] * mm[(e, BASE)] for e in ee) / t_base * 100
    visibles, omitidos = [], []
    for dim in ['region', 'provincia', 'cadena', 'tipo']:
        for g in sorted({g for (d, g) in meses_g if d == dim}):
            n = len(sucs_g.get((dim, g, ultimo()), ()))
            if n >= MIN_SUC:
                i, v, _ = serie(dim, g, eans_all)
                res[(dim, g)] = (i, v)
                visibles.append((dim, g, n))
            else:
                omitidos.append(f'{g} ({n})')
    return res, pond, visibles, omitidos


def verificar(env, xls, etiqueta, maestro_cambiado=False):
    print(f'\n-- {etiqueta} --')
    dl = env['detalle_long']
    chequear(dl is not None and len(dl) > 0, 'detalle_long armado')
    res, pond, visibles, omitidos = esperado(maestro_cambiado)
    per = meses_vis() + PER_W
    chequear(sorted(dl['periodo'].unique()) == sorted(per),
             f'periodos: {len(meses_vis())} meses + {len(PER_W)} semanas ({sorted(dl["periodo"].unique())})')

    def fila(dim, g):
        g = 'Total Media' if (dim, g) == ('nacional', 'Total') else g
        return dl[(dl['seccion'] == dim) & (dl['fila'] == g)].set_index('periodo')

    malos = []
    for (dim, g), (idx, var) in res.items():
        f = fila(dim, g)
        if len(f) != len(per):
            malos.append(f'{dim}/{g}: {len(f)} periodos')
            continue
        for p in per:
            if not (cerca(f.loc[p, 'indice'], idx[p]) and cerca(f.loc[p, 'variacion_%'], var.get(p, np.nan))):
                malos.append(f'{dim}/{g}/{p}: {f.loc[p, "indice"]:.6f} vs {idx[p]:.6f}, '
                             f'{f.loc[p, "variacion_%"]:.6f} vs {var.get(p, np.nan):.6f}')
    chequear(not malos, f'{len(res)} filas (Total, rubros, categorias, regiones, provincias, cadenas, tipos) '
                        f'= calculo independiente en indice y variacion' + (f': {malos[:4]}' if malos else ''))

    sn = env['serie_nac_dict']['cantidad_03'].set_index('mes')
    t = fila('nacional', 'Total')
    ok = all(cerca(t.loc[p, 'variacion_%'], sn.loc[p, 'variacion_mensual_%']) for p in meses_vis())
    ok &= all(cerca(t.loc[p, 'indice'], sn.loc[p, 'canasta_nacional_ponderada'] /
                    sn.loc[BASE, 'canasta_nacional_ponderada'] * 100) for p in meses_vis())
    chequear(ok, f'Total = serie oficial de la CELDA 9 (indice base {BASE} y variacion)')

    pd_ok = all(cerca(fila(d, g)['pond'].iloc[0], pond[(d, g)]) for (d, g) in pond)
    chequear(pd_ok, 'Pond de rubros y categorias = participacion en el costo del mes base')
    rub = [g for (d, g) in pond if d == 'rubro']
    suma = all(cerca(sum(pond[('rubro', g)] * fila('rubro', g).loc[p, 'indice'] for g in rub) / 100,
                     t.loc[p, 'indice']) for p in meses_vis())
    chequear(suma, 'Total = suma de Pond x indice de los rubros / 100')
    vis = dl[dl['seccion'].isin(['region', 'provincia', 'cadena', 'tipo'])].drop_duplicates(['seccion', 'fila'])
    chequear(sorted(zip(vis['seccion'], vis['fila'])) == sorted((d, g) for d, g, _ in visibles),
             f'grupos visibles con >= {MIN_SUC} sucursales: {len(visibles)}')
    suc_ok = all(int(fila(d, g)['sucursales'].iloc[0]) == n for d, g, n in visibles)
    chequear(suc_ok, 'columna Sucursales = sucursales del grupo en el ultimo mes completo')

    # ── La hoja del Excel ──
    wb = load_workbook(xls)
    ws = wb.worksheets[0]
    chequear(ws.title == 'analisis_detallado', f'primera hoja: {ws.title}')
    chequear([w.title for w in wb.worksheets if w.sheet_view.tabSelected] == ['analisis_detallado']
             and wb.active.title == 'analisis_detallado', 'unica pestaña seleccionada y activa')
    chequear({'Evolucion_IPC', 'Serie_precios', 'Productos', 'Cadenas'} <= set(wb.sheetnames)
             if env['DESAGREGAR_CANASTAS'] else {'Evolucion_IPC', 'Serie_precios'} <= set(wb.sheetnames),
             'las demas hojas siguen en el Excel')
    filas_a = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'Total Media']
    hdrs = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == f'Pond. {ABR[int(BASE[5:])]}-26 %']
    ok_x = len(filas_a) == 2 and len(hdrs) == 2
    lab = {}
    if ok_x:
        for c in range(1, ws.max_column + 1):
            if ws.cell(hdrs[0], c).value:
                lab[ws.cell(hdrs[0], c).value] = c
        esperadas = etiquetas()
        ok_x &= all(x in lab for x in esperadas) and len(lab) == 4 + len(esperadas)
        for (p, x) in zip(meses_vis() + PER_W, esperadas):
            ok_x &= cerca(ws.cell(filas_a[0], lab[x]).value, t.loc[p, 'indice'], 1e-5)
            vv = ws.cell(filas_a[1], lab[x]).value
            ok_x &= cerca(np.nan if vv is None else vv, t.loc[p, 'variacion_%'], 1e-5)
    chequear(ok_x, f'celdas del Excel = detalle_long (bloques indice y variacion, columnas {list(lab)[:6]}...)')
    notas = ' '.join(str(ws.cell(r, 1).value or '') for r in range(1, ws.max_row + 1))
    chequear(all(o in notas for o in omitidos) and 'Cadena 99 (1)' in omitidos,
             f'grupos omitidos en las notas: {omitidos}')
    return dl


if __name__ == '__main__':
    tmp = pathlib.Path(tempfile.mkdtemp(prefix='test_detalle_nb02_'))
    (tmp / 'sepa').mkdir()
    escribir_zips(tmp / 'sepa')
    cache = tmp / 'cache'
    cache.mkdir()

    coto_abr = [1 for si, e, fe, p in OBS if SUCS[si][0] == '12' and e == E_ACEITE and fe // 100 == 202604]
    coto_may = [1 for si, e, fe, p in OBS if SUCS[si][0] == '12' and e == E_ACEITE and fe // 100 == 202605]
    chequear(not coto_abr and coto_may, 'escenario: Coto sin aceite en abril y con aceite en mayo')
    env1, xls1, lec1 = correr(tmp, 'r1', ['cantidad_03'], cache, contar_lecturas=True)
    dl1 = verificar(env1, xls1, 'primera corrida (sin cache)')
    chequear(lec1 == ['2026-03', '2026-04', '2026-05'],
             f'lee del SEPA los meses viejos de la ventana (jun y jul vienen de la 19b): {lec1}')
    # Re-correr la 19c sin la 19b: avisa y no rompe (las observaciones ya se liberaron)
    exec(compile(C19C, '<CELDA 19c>', 'exec'), env1)
    chequear(env1['detalle_long'] is None and load_workbook(xls1).sheetnames[0] == 'analisis_detallado',
             'la 19c sola (sin la 19b antes) avisa, no rompe y deja la hoja anterior')
    env_x = {'__builtins__': __builtins__, 'DETALLE_CANASTA': 'cantidad_03', 'CANASTAS': {'cantidad_03': {}}}
    exec(compile(C19C, '<CELDA 19c>', 'exec'), env_x)
    chequear(env_x['detalle_long'] is None, 'la 19c en una sesion sin la 19b (sin _DET ni _OBS) avisa y no rompe')
    cfiles = list(cache.glob('detalle_*_v1.parquet'))
    meses_c = sorted(pd.read_parquet(cfiles[0])['mes'].unique()) if cfiles else []
    chequear(len(cfiles) == 1 and meses_c == ['2026-03', '2026-04', '2026-05', '2026-06'],
             f'cache con los meses cerrados: {meses_c}')

    env2, xls2, lec2 = correr(tmp, 'r2', [], cache, contar_lecturas=True)
    dl2 = verificar(env2, xls2, 'segunda corrida (cache) y DESAGREGAR_CANASTAS vacio')
    chequear(lec2 == [], f'con el cache no relee meses viejos: {lec2}')
    m = dl1.merge(dl2, on=['seccion', 'fila', 'periodo'], suffixes=('_1', '_2'))
    chequear(len(m) == len(dl1) and np.allclose(m['indice_1'], m['indice_2'], equal_nan=True)
             and np.allclose(m['variacion_%_1'], m['variacion_%_2'], equal_nan=True),
             'segunda corrida = primera')

    env3, xls3, lec3 = correr(tmp, 'r3', ['cantidad_03'], cache, contar_lecturas=True, maestro_cambiado=True)
    verificar(env3, xls3, 'maestro cambiado (una sucursal de AMBA pasa a Pampeana)', maestro_cambiado=True)
    chequear(lec3 == ['2026-03', '2026-04', '2026-05'] and len(list(cache.glob('detalle_*_v1.parquet'))) == 2,
             f'maestro distinto -> clave nueva -> relee: {lec3}')

    # Julio COMPLETO: el ultimo mes completo es el mes en curso (sin columna parcial, 4 + 4 semanas)
    COMPLETO = True
    DIAS[202607] = 31
    COLS_M, BASE = ['2026-05', '2026-06', '2026-07'], '2026-04'
    PER_W = [f'2026-0{m}-S{k}' for m in (6, 7) for k in range(1, 5)]
    OBS = generar()
    (tmp / 'sepa_c').mkdir()
    escribir_zips(tmp / 'sepa_c')
    cache_c = tmp / 'cache_c'
    cache_c.mkdir()
    env4, xls4, lec4 = correr(tmp, 'r4', ['cantidad_03'], cache_c, contar_lecturas=True, sepa='sepa_c')
    verificar(env4, xls4, 'julio completo (sin mes parcial)')
    meses_c = sorted(pd.read_parquet(next(cache_c.glob('detalle_*_v1.parquet')))['mes'].unique())
    chequear(lec4 == ['2026-04', '2026-05'] and meses_c == ['2026-04', '2026-05', '2026-06'],
             f'lee abr y may; el cache no guarda el mes en curso: {lec4} / {meses_c}')
    shutil.rmtree(tmp, ignore_errors=True)
    print('\nRESULTADO:', 'OK' if not fallas else f'FALLA ({len(fallas)})')
    raise SystemExit(1 if fallas else 0)
